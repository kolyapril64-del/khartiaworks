"""Генератор робочої книги «Добова доповідь ППО» (Excel, .xlsx).

Підрозділ вставляє готові тексти доповідей у журнал, а поля журналу
(час, тип цілі, район, квадрат, засіб, результат, витрати) розбираються
з тексту формулами за таблицею ключових слів на аркуші «Довідники».
Аркуші «Зведення», «За типами цілей» і «Розрахунок» рахуються автоматично.
Аркуш «Картки» складає текст картки об'єкта з перерахунком координат MGRS ↔ WGS84.

Використання:
    python build_workbook.py config.json output.xlsx

Формат config.json див. у config.example.json.
"""

import json
import math
import os
import sys
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------------------------------------------------------------- розміри
UNIT_SLOTS = 15      # скільки підрозділів вміщує зведення
TYPE_SLOTS = 30      # типів цілей у довіднику
MEANS_SLOTS = 150    # засобів ураження (видів озброєння)
AMMO_SLOTS = 200     # пар «засіб — боєприпас»
RESULT_SLOTS = 10
STATUS_SLOTS = 8
STRIKE_SLOTS = 5
ENEMY_SLOTS = 14     # засобів противника (стовпці у зведенні ударів)
RADAR_SLOTS = 8
OBJECT_SLOTS = 20    # типів об'єктів для карток
CARD_SLOTS = 30      # карток (стовпців) на аркуші «Картки»

# Аркуш підрозділу: фіксовані адреси, на які посилаються зведення і розрахунок
CREW_SLOTS = 20      # рядків у таблиці екіпажів перехоплювачів
CREW_ROWS = range(19, 19 + CREW_SLOTS)
CREW_TOTAL = CREW_ROWS[-1] + 1
J_HEAD = CREW_TOTAL + 3                                # журнал цілей: заголовок, приклад, дані
J_EXAMPLE, J_FIRST = J_HEAD + 1, J_HEAD + 2
J_LAST = J_FIRST + 299
S_LAST = J_FIRST + 99                                  # удари противника
JC = dict(num="A", text="B", time="C", type="D", place="E", grid="F", means="G", ammo="H",
          crew="I", result="J", spent="K", lost="L", note="M", group="N", check="O")
SC = dict(num="Q", text="R", time="S", kind="T", enemy="U", qty="V", target="W", losses="X")
U_NAME, U_DATE, U_DUTY, U_FROM, U_TO = "B2", "E2", "B3", "E3", "F3"
RADAR_ROWS = range(11, 15)
RADAR_TOTAL = 15

# Службові значення (на них спираються формули — не перейменовувати)
DESTROYED = "Знищено"
LOST = "Борт втрачено"
NOT_FOUND = "Ціль не виявлено, борт повернули"
PATROL = "Без цілі (патрулювання)"
G_FPV, G_SMALL, G_AAA, G_SAM, G_OTHER = "ФПВ", "Стр.зб", "ЗА", "ЗРК/ПЗРК", "Інші"
GROUPS = [(G_FPV, "FPV-перехоплювачі"), (G_SMALL, "Стрілецька зброя"), (G_AAA, "Зенітна артилерія"),
          (G_SAM, "ЗРК / ПЗРК"), (G_OTHER, "Інші засоби")]
READY = "боєготові"
STRIKE_CODES = [("РУ", "ракетний удар"), ("АУ", "авіаційний удар"), ("УДК", "удар дроном-камікадзе")]

DEFAULT_TARGETS = [
    "Shahed", "Гербера", "Ланцет", "Молнія", "FPV", "FPV (оптоволокно)", "Бомбер",
    "Mavic / Autel", "DJI Matrice", "Орлан-10/30", "Supercam", "Скат", "Zala", "V2U",
    "НВТ", "Італмас", "КВО (Князь Віщий Олег)", "Сокіл", "Кощей", "Merlin", "Куб",
    "Легіонер", "Аеростат", "Гексокоптер", PATROL,
]
# (вид озброєння, боєприпас, група); config.json може задати власний перелік у "weapons"
DEFAULT_WEAPONS = [
    ("FPV-перехоплювач", "", G_FPV), ("Стрілецька зброя", "Набої", G_SMALL),
    ("Зенітна артилерія", "", G_AAA), ("ЗРК", "", G_SAM), ("ПЗРК", "", G_SAM),
    ("Скид котушки", "", G_OTHER), ("Переріз оптоволокна", "", G_OTHER), ("РЕБ", "—", G_OTHER),
    ("Інше", "", G_OTHER),
]
DEFAULT_RESULTS = [DESTROYED, "Пошкоджено", "Не знищено", NOT_FOUND, LOST, "Патрулювання"]
DEFAULT_STATUS = [READY, "обмежено боєготові", "відновлення", "ремонт", "ротація"]
DEFAULT_OBJECTS = ["РЛС", "ЗРК", "ПЗРК", "Зенітна установка", "Мобільна вогнева група",
                   "Екіпаж FPV-перехоплювачів", "Пост спостереження", "Засіб РЕБ", "Пункт управління"]
DEFAULT_READINESS = ["боєздатний", "обмежено боєздатний", "небоєздатний", "знищений"]
DEFAULT_AFFILIATION = ["Дружній", "Ворожий", "Нейтральний", "Невідомий"]
DEFAULT_ENEMY = ["Шахед", "Гербера", "Герань", "Бандероль", "Ланцет", "Молнія", "Італмас", "Куб", "FPV",
                 "КАБ", "КАР", "Ракета", "Інше"]

# Ключові слова для розбору тексту: (фрагмент тексту, значення).
# Регістр не важливий; якщо збіглося кілька слів — береться найдовше,
# за однакової довжини — нижче у списку.
KW_TARGETS = [
    ("shahed", "Shahed"), ("шахед", "Shahed"), ("герань", "Shahed"), ("гербер", "Гербера"),
    ("gerbera", "Гербера"), ("ланцет", "Ланцет"), ("lancet", "Ланцет"), ("молні", "Молнія"),
    ("молния", "Молнія"), ("fpv", "FPV"), ("фпв", "FPV"), ("оптоволок", "FPV (оптоволокно)"),
    ("бомбер", "Бомбер"), ("mavic", "Mavic / Autel"), ("мавік", "Mavic / Autel"),
    ("autel", "Mavic / Autel"), ("matrice", "DJI Matrice"), ("орлан", "Орлан-10/30"),
    ("orlan", "Орлан-10/30"), ("supercam", "Supercam"), ("суперкам", "Supercam"), ("скат", "Скат"),
    ("zala", "Zala"), ("зала", "Zala"), ("v2u", "V2U"), ("нвт", "НВТ"), ("італмас", "Італмас"),
    ("ілтамас", "Італмас"), ("кво", "КВО (Князь Віщий Олег)"), ("віщий олег", "КВО (Князь Віщий Олег)"),
    ("сокіл", "Сокіл"), ("сокол", "Сокіл"), ("кощей", "Кощей"), ("кащей", "Кощей"),
    ("merlin", "Merlin"), ("мерлін", "Merlin"), ("куб", "Куб"), ("kub", "Куб"),
    ("легіонер", "Легіонер"), ("legioner", "Легіонер"), ("аеростат", "Аеростат"),
    ("гексокоптер", "Гексокоптер"), ("гекса", "Гексокоптер"),
]
# Пріоритет: 0 — запасне (загальне) слово, порожньо = 1, більше — важливіше
KW_MEANS = [
    ("перехоплювач", "FPV-перехоплювач", 0), ("стрілецьк", "Стрілецька зброя", 0),
    ("стр.зб", "Стрілецька зброя", 0), ("кулемет", "Стрілецька зброя", 0), ("обстріл", "Стрілецька зброя", 0),
    ("ак-74", "Стрілецька зброя"), ("акс-74", "Стрілецька зброя"), ("пкм", "Стрілецька зброя"),
    ("котушк", "Скид котушки"),
    ("катушк", "Скид котушки"), ("переріз", "Переріз оптоволокна"), ("пзрк", "ПЗРК", 0),
    ("зрк", "ЗРК", 0), ("реб", "РЕБ"),
]
KW_RESULTS = [
    ("знищено", DESTROYED), ("збито", DESTROYED), (" не знищено", "Не знищено"),
    ("пошкоджено", "Пошкоджено"), (" не виявлено", NOT_FOUND), ("борт повернули", NOT_FOUND),
    ("борт повернувся", NOT_FOUND), ("борт втрачено", LOST), ("втрачено борт", LOST),
    ("невлучання", "Не знищено"), ("результат: втрата", LOST),
]
KW_ENEMY = [
    ("шахед", "Шахед"), ("shahed", "Шахед"), ("геран", "Герань"), ("гербер", "Гербера"),
    ("gerbera", "Гербера"), ("бандерол", "Бандероль"), ("ланцет", "Ланцет"), ("молні", "Молнія"),
    ("італмас", "Італмас"), ("куб", "Куб"), ("fpv", "FPV"), ("фпв", "FPV"), ("каб", "КАБ"),
    (" кар ", "КАР"), (" кар(", "КАР"), ("авіаційна ракета", "КАР"), ("ракет", "Ракета"),
]
KW_STRIKE = [("ракет", "РУ"), ("авіац", "АУ"), ("каб", "АУ"), ("фаб", "АУ"), ("умпк", "АУ"),
             (" кар ", "АУ"), (" кар(", "АУ"), ("авіаційна ракета", "АУ")]

# «Довідники»: стовпці списків (значення, супутній стовпець, кількість рядків)
RL = {"units": ("A", "B", UNIT_SLOTS), "targets": ("D", None, TYPE_SLOTS),
      "means": ("F", "G", MEANS_SLOTS), "ammo": ("I", "J", AMMO_SLOTS),
      "results": ("L", None, RESULT_SLOTS), "status": ("N", None, STATUS_SLOTS),
      "strike": ("P", "Q", STRIKE_SLOTS), "enemy": ("S", None, ENEMY_SLOTS),
      "radar": ("U", None, RADAR_SLOTS), "groups": ("W", None, len(GROUPS)),
      "objects": ("AS", None, OBJECT_SLOTS), "readiness": ("AU", None, 8), "affil": ("AW", None, 8)}
# Таблиці ключових слів: слово, значення, пріоритет, кількість рядків
KW = {"targets": ("Y", "Z", "AA", 60), "means": ("AC", "AD", "AE", 250),
      "results": ("AG", "AH", "AI", 20), "enemy": ("AK", "AL", "AM", 25),
      "strike": ("AO", "AP", "AQ", 10)}
NOT_FOUND_ROW, FALLBACK_ROW = 999, 998   # порожній рядок / рядок значення «за замовчуванням»

# ---------------------------------------------------------------- стиль
FONT = "Arial"
C_NAVY, C_TEAL, C_SLATE = "1F3A5F", "0F6E73", "44546A"
C_INPUT, C_PARSED, C_AUTO, C_TILE, C_EXAMPLE = "FFF7D6", "E6F0FA", "EEF0F3", "F3F6FA", "F7F7F7"
C_LINE = "C9CED6"

THIN = Side(style="thin", color=C_LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = '0;-0;"–"'
PCT = '0%;-0%;"–"'


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def font(size=10, bold=False, color="1A1A1A", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def q(sheet):
    """Назва аркуша в лапках для посилань."""
    return "'" + sheet.replace("'", "''") + "'"


def cells(ws, rng):
    for row in ws[rng] if ":" in rng else [[ws[rng]]]:
        for c in row:
            yield c


def style(ws, rng, *, f=None, fnt=None, border=BOX, align=None, fmt=None, unlocked=False, merge=False):
    for c in cells(ws, rng):
        if f:
            c.fill = f
        if fnt:
            c.font = fnt
        if border:
            c.border = border
        if align:
            c.alignment = align
        if fmt:
            c.number_format = fmt
        if unlocked:
            c.protection = Protection(locked=False)
    if merge and ":" in rng:
        ws.merge_cells(rng)


CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def title(ws, rng, text, size=14):
    ws[rng.split(":")[0]] = text
    style(ws, rng, f=fill(C_NAVY), fnt=font(size, True, "FFFFFF"), border=None,
          align=Alignment(horizontal="left", vertical="center", indent=1), merge=True)


def band(ws, rng):
    style(ws, rng, f=fill(C_NAVY), border=None, merge=":" in rng)


def section(ws, rng, text, color=C_TEAL):
    ws[rng.split(":")[0]] = text
    style(ws, rng, f=fill(color), fnt=font(10, True, "FFFFFF"), border=None,
          align=Alignment(horizontal="left", vertical="center", indent=1), merge=":" in rng)


def header(ws, rng, text=None):
    if text is not None:
        ws[rng.split(":")[0]] = text
    style(ws, rng, f=fill(C_SLATE), fnt=font(9, True, "FFFFFF"), align=CENTER, merge=":" in rng)


def label(ws, rng, text):
    ws[rng.split(":")[0]] = text
    style(ws, rng, f=fill(C_AUTO), fnt=font(9, True, C_SLATE), align=LEFT, merge=":" in rng)


def inp(ws, rng, value=None, fmt=None, align=LEFT):
    if value is not None:
        ws[rng.split(":")[0]] = value
    style(ws, rng, f=fill(C_INPUT), fnt=font(), align=align, fmt=fmt, unlocked=True, merge=":" in rng)


def parsed(ws, rng, formula, fmt=None, align=LEFT):
    """Розібране з тексту поле: формула, але клітинку можна виправити вручну."""
    ws[rng] = formula
    style(ws, rng, f=fill(C_PARSED), fnt=font(), align=align, fmt=fmt, unlocked=True)


def auto(ws, rng, value=None, fmt=None, bold=False, align=CENTER):
    if value is not None:
        ws[rng.split(":")[0]] = value
    style(ws, rng, f=fill(C_AUTO), fnt=font(10, bold), align=align, fmt=fmt, merge=":" in rng)


def tile(ws, lab_rng, val_rng, text, formula, fmt=NUM):
    ws[lab_rng.split(":")[0]] = text
    style(ws, lab_rng, f=fill(C_TILE), fnt=font(8, True, "5B6B7F"), align=CENTER,
          border=Border(left=THIN, right=THIN, top=THIN), merge=":" in lab_rng)
    ws[val_rng.split(":")[0]] = formula
    style(ws, val_rng, f=fill(C_TILE), fnt=font(18, True, C_NAVY), align=CENTER, fmt=fmt,
          border=Border(left=THIN, right=THIN, bottom=THIN), merge=":" in val_rng)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def list_dv(ws, name, ranges, prompt=None, strict=True):
    dv = DataValidation(type="list", formula1=name, allow_blank=True, showDropDown=False,
                        errorStyle="stop" if strict else "warning")
    dv.error, dv.errorTitle = "Оберіть значення зі списку (Довідники).", "Невірне значення"
    if prompt:
        dv.prompt, dv.showInputMessage = prompt, True
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)


def num_dv(ws, ranges, hi=999):
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2=str(hi), allow_blank=True)
    dv.error, dv.errorTitle = f"Ціле число від 0 до {hi}.", "Невірне значення"
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)


def time_dv(ws, ranges):
    dv = DataValidation(type="time", operator="between", formula1="0", formula2="0.999305555555556",
                        allow_blank=True, errorStyle="warning")
    dv.error, dv.errorTitle = "Вводьте час у форматі ГГ:ХХ, напр. 07:24.", "Формат часу"
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)


def date_dv(ws, ranges):
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        errorStyle="warning")
    dv.error, dv.errorTitle = "Вводьте дату у форматі ДД.ММ.РРРР.", "Формат дати"
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)


def protect(ws):
    p = ws.protection
    p.sheet = True
    p.autoFilter = False      # фільтри дозволені
    p.formatColumns = False   # ширину стовпців можна змінювати
    p.formatRows = False      # і висоту рядків


def page(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False


def hhmm(ref):
    return f'RIGHT("0"&HOUR({ref}),2)&":"&RIGHT("0"&MINUTE({ref}),2)'


def emblem(ws, path, col, row, box_w, box_h, off_x=0, off_y=0):
    """Вставляє шеврон, вписаний у рамку box_w×box_h пікселів, від клітинки (col, row), 0-based."""
    if not path or not os.path.exists(path):
        return
    from PIL import Image as PILImage, ImageChops

    im = PILImage.open(path).convert("RGBA")
    alpha = im.getchannel("A")
    if alpha.getextrema()[0] < 255:
        box = alpha.point(lambda a: 255 if a > 8 else 0).getbbox()
    else:   # непрозорий фон (JPG): обрізаємо поля кольору кута
        bg = PILImage.new("RGBA", im.size, im.getpixel((0, 0)))
        box = ImageChops.difference(im, bg).convert("L").point(lambda v: 255 if v > 24 else 0).getbbox()
    if box:
        im = im.crop(box)
    scale = min(box_w / im.width, box_h / im.height)
    dw, dh = max(1, round(im.width * scale)), max(1, round(im.height * scale))
    im = im.resize((dw * 3, dh * 3), PILImage.LANCZOS)
    buf = BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    img = XLImage(buf)
    img.width, img.height = dw, dh
    ox = off_x + (box_w - dw) // 2
    img.anchor = OneCellAnchor(
        _from=AnchorMarker(col=col, colOff=pixels_to_EMU(ox), row=row, rowOff=pixels_to_EMU(off_y)),
        ext=XDRPositiveSize2D(pixels_to_EMU(dw), pixels_to_EMU(dh)))
    ws.add_image(img)


# ================================================================ Довідники
REF = "Довідники"


def ref_range(col, n):
    return f"{q(REF)}!${col}$5:${col}${4 + n}"


def rl(kind, second=False):
    col, col2, n = RL[kind]
    return ref_range(col2 if second else col, n)


def kw_match(kind, text, fallback=str(NOT_FOUND_ROW)):
    """Значення найкращого ключового слова, знайденого в text ("" — нічого не знайдено).

    Бал = (пріоритет*100 + довжина слова)*1000 + рядок; MOD(бал,1000) — рядок значення.
    Якщо нічого не знайдено, береться рядок fallback (порожній або зі значенням за замовчуванням).
    """
    a, c, p, n = KW[kind]
    A, P = ref_range(a, n), ref_range(p, n)
    score = f'(({P}+({P}=""))*100+LEN({A}))*1000+ROW({A})'
    v = f'SUMPRODUCT(MAX(ISNUMBER(SEARCH({A},{text}))*({A}<>"")*{score},{fallback}))'
    return f'INDEX({q(REF)}!${c}:${c},MOD({v},1000))&""'


def build_reference(wb, cfg):
    ws = wb.create_sheet(REF)
    ws.sheet_properties.tabColor = "8C96A3"
    page(ws)
    title(ws, "A1:AW1", "ДОВІДНИКИ — списки для меню і ключові слова для розбору доповідей")
    label(ws, "A2", "Назва об'єднання")
    inp(ws, "B2:D2", cfg["org_name"])
    ws["F2"] = ("Нові значення дописуйте в порожні жовті клітинки. Виділені жирним — службові: на них "
                "спираються формули, не змінюйте їх. Боєприпаси одного засобу тримайте підряд.")
    style(ws, "F2:W2", fnt=font(9, italic=True, color="5B6B7F"), border=None, align=LEFT, merge=True)
    ws["Y2"] = ("Розбір тексту: шукається кожне слово (регістр не важливий). Перемагає більший пріоритет "
                "(порожньо = 1, 0 = запасне загальне слово), далі довше слово. Тип цілі шукається лише у "
                "фразі після «виявлено».")
    style(ws, "Y2:AQ2", fnt=font(9, italic=True, color="5B6B7F"), border=None, align=LEFT, merge=True)
    ws.row_dimensions[2].height = 42
    section(ws, "A3:W3", "СПИСКИ ДЛЯ ВИПАДАЮЧИХ МЕНЮ")
    section(ws, "Y3:AQ3", "КЛЮЧОВІ СЛОВА ДЛЯ РОЗБОРУ ТЕКСТУ ДОПОВІДЕЙ", color=C_SLATE)
    section(ws, "AS3:AW3", "СПИСКИ ДЛЯ КАРТОК")

    def column(col, head, values, slots, service=(), extra=()):
        """extra: [(стовпець, заголовок, значення, список_для_меню | None, числовий?)]"""
        header(ws, f"{col}4", head)
        for xcol, xhead, *_ in extra:
            header(ws, f"{xcol}4", xhead)
        for i in range(slots):
            c = f"{col}{5 + i}"
            v = values[i] if i < len(values) else None
            svc = v in service
            if svc:
                auto(ws, c, v, bold=True, align=LEFT)
                ws[c].comment = Comment("Службове значення — використовується у формулах.", "ППО")
            else:
                inp(ws, c, v)
            for xcol, _, xvals, _, numeric in extra:
                xv = xvals[i] if i < len(xvals) else None
                (auto if svc else inp)(ws, f"{xcol}{5 + i}", xv, align=CENTER if numeric else LEFT)
        for xcol, _, _, dv, numeric in extra:
            rng = f"{xcol}5:{xcol}{4 + slots}"
            if dv:
                list_dv(ws, dv, [rng])
            if numeric:
                num_dv(ws, [rng], hi=9)

    weapons = [tuple(w) for w in cfg.get("weapons", DEFAULT_WEAPONS)]
    means, groups = [], {}
    for name, _, grp in weapons:
        if name not in groups:
            means.append(name)
            groups[name] = grp
    ammo = [(name, a) for name, a, _ in weapons if a]
    if len(means) > MEANS_SLOTS or len(ammo) > AMMO_SLOTS:
        raise SystemExit("Забагато засобів або боєприпасів для довідника")

    units = cfg["units"]
    column("A", "Код аркуша", [u["code"] for u in units], UNIT_SLOTS,
           extra=[("B", "Повна назва підрозділу", [u["name"] for u in units], None, False)])
    column("D", "Типи цілей", DEFAULT_TARGETS, TYPE_SLOTS, service=(PATROL,))
    column("F", "Засіб ураження (вид озброєння)", means, MEANS_SLOTS,
           extra=[("G", "Група", [groups[m] for m in means], "lst_groups", False)])
    column("I", "Засіб ураження", [m for m, _ in ammo], AMMO_SLOTS,
           extra=[("J", "Боєприпас", [a for _, a in ammo], None, False)])
    list_dv(ws, "lst_means", [f"I5:I{4 + AMMO_SLOTS}"])
    column("L", "Результати", DEFAULT_RESULTS, RESULT_SLOTS, service=(DESTROYED, LOST))
    column("N", "Стан екіпажу", DEFAULT_STATUS, STATUS_SLOTS, service=(READY,))
    column("P", "Тип удару", [c for c, _ in STRIKE_CODES], STRIKE_SLOTS,
           service=tuple(c for c, _ in STRIKE_CODES),
           extra=[("Q", "Розшифровка", [d for _, d in STRIKE_CODES], None, False)])
    column("S", "Засоби противника", DEFAULT_ENEMY, ENEMY_SLOTS)
    column("U", "Типи РЛС", cfg.get("radars", []), RADAR_SLOTS)
    column("W", "Групи засобів", [g for g, _ in GROUPS], len(GROUPS), service=tuple(g for g, _ in GROUPS))
    column("AS", "Типи об'єктів", DEFAULT_OBJECTS, OBJECT_SLOTS)
    column("AU", "Боєздатність", DEFAULT_READINESS, RL["readiness"][2])
    column("AW", "Приналежність", DEFAULT_AFFILIATION, RL["affil"][2])

    for kind, name in [("units", "lst_units"), ("targets", "lst_targets"), ("means", "lst_means"),
                       ("groups", "lst_groups"), ("results", "lst_results"), ("status", "lst_status"),
                       ("strike", "lst_strike"), ("enemy", "lst_enemy"), ("radar", "lst_radar"),
                       ("objects", "lst_objects"), ("readiness", "lst_readiness"), ("affil", "lst_affil")]:
        wb.defined_names[name] = DefinedName(name, attr_text=rl(kind))
    wb.defined_names["lst_unit_names"] = DefinedName("lst_unit_names", attr_text=rl("units", True))

    def kw_rows(pairs):
        pairs = [tuple(x) for x in pairs]
        return ([x[0] for x in pairs], [x[1] for x in pairs],
                [x[2] if len(x) > 2 else None for x in pairs])

    for kind, head, pairs, dv in [("targets", "→ Тип цілі", KW_TARGETS, "lst_targets"),
                                  ("means", "→ Засіб ураження", cfg.get("means_keywords", KW_MEANS), "lst_means"),
                                  ("results", "→ Результат", KW_RESULTS, "lst_results"),
                                  ("enemy", "→ Засіб противника", KW_ENEMY, "lst_enemy"),
                                  ("strike", "→ Тип удару", KW_STRIKE, "lst_strike")]:
        a, c, p, n = KW[kind]
        words, vals, prios = kw_rows(pairs)
        if len(words) > n:
            raise SystemExit(f"Забагато ключових слів ({kind})")
        column(a, "Слово в тексті", words, n,
               extra=[(c, head, vals, dv, False), (p, "Пріор.", prios, None, True)])

    # значення «за замовчуванням», якщо слово не знайдено (рядок FALLBACK_ROW)
    for kind, val, note in [("targets", PATROL, "якщо в тексті «патрулювання» і тип не знайдено"),
                            ("strike", "УДК", "якщо тип удару не знайдено")]:
        a, c, _, _ = KW[kind]
        ws[f"{a}{FALLBACK_ROW}"] = "службове: " + note
        ws[f"{a}{FALLBACK_ROW}"].font = font(8, italic=True, color="7F7F7F")
        auto(ws, f"{c}{FALLBACK_ROW}", val, bold=True, align=LEFT)

    w = {"A": 12, "B": 30, "D": 26, "F": 30, "G": 11, "I": 30, "J": 22, "L": 30, "N": 20, "P": 10,
         "Q": 22, "S": 20, "U": 20, "W": 14, "AS": 26, "AU": 22, "AW": 16}
    for a, c, p, _ in KW.values():
        w.update({a: 18, c: 26, p: 7})
    for i in range(1, 50):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = w.get(col, 2)
    ws.freeze_panes = "A5"
    protect(ws)
    return ws


# ================================================================ аркуш підрозділу
REGION_WORDS = []   # початок назви області в доповідях (з config.json "region_words"), напр. ["Київ"]


def parse_formulas(r):
    """Формули розбору тексту доповіді в рядку r журналу цілей."""
    B = f"{JC['text']}{r}"

    def time_of(t):
        return (f'IFERROR(TIMEVALUE(SUBSTITUTE(LEFT(TRIM({t}),5),".",":")),'
                f'IFERROR(TIMEVALUE(SUBSTITUTE(LEFT(TRIM({t}),4),".",":")),""))')

    # тип цілі — лише у фразі від «виявлено» (або «БпЛА») до крапки
    s = f'IFERROR(SEARCH("виявлено",{B}),IFERROR(SEARCH("ціль:",{B}),IFERROR(SEARCH("бпла",{B}),0)))'
    seg = f'MID({B},{s},IFERROR(MIN(SEARCH(".",{B},{s})-{s},70),70))'
    patrol = f'IF(ISNUMBER(SEARCH("патрулюван",{B})),{FALLBACK_ROW},{NOT_FOUND_ROW})'
    G = f"{JC['means']}{r}"
    AW, AA = rl("ammo"), rl("ammo", True)
    ammo_score = (f'SUMPRODUCT(MAX(({AW}={G})*ISNUMBER(SEARCH({AA},{B}))*({AA}<>"")'
                  f'*(LEN({AA})*1000+ROW({AA})),{NOT_FOUND_ROW}))')
    ammo = (f'IF({G}="","",IF(COUNTIF({AW},{G})=1,INDEX({AA},MATCH({G},{AW},0)),'
            f'INDEX({q(REF)}!${RL["ammo"][1]}:${RL["ammo"][1]},MOD({ammo_score},1000))&""))')

    # район: після «н.п.» / «в районі» / «в р-ні» до коми, назви області, дужки чи крапки
    m = (f'IFERROR(SEARCH("н.п.",{B})+4,IFERROR(SEARCH("в районі",{B})+8,'
         f'IFERROR(SEARCH("в р-ні",{B})+6,0)))')
    rest = f"MID({B},{m},50)"
    stops = [f" {w}" for w in REGION_WORDS] + [",", " (", ". ", " виявлено", " обл", " кв"]
    end = "MIN(" + ",".join(f'IFERROR(SEARCH("{x}",{rest}),51)' for x in stops) + ")"
    after_marker = f"TRIM(LEFT({rest},{end}-1))"
    # без маркера: останній фрагмент через кому перед назвою області, «кв.» чи «орієнтир»
    cuts = [x for w in REGION_WORDS for x in (f", {w}", f" {w}")] + [", кв.", " кв.", ", орієнтир"]
    cut = "MIN(" + ",".join(f'IFERROR(SEARCH("{x}",{B}),999)' for x in cuts) + ")"
    before_region = (f'IF({cut}=999,"",TRIM(RIGHT(SUBSTITUTE(LEFT({B},{cut}-1),",",'
                     f'REPT(" ",100)),100)))')

    # квадрат MGRS: «(37U XX 12345 67890)»
    p = f'SEARCH("(3?U ",{B})'
    grid = f'IFERROR(MID({B},{p}+1,SEARCH(")",{B},{p})-{p}-1),IFERROR(TRIM(MID({B},SEARCH("3?U ??",{B}),18)),""))'

    # екіпаж: «ЗПМ «Назва»», «...) “Назва” пілот», «Екіпаж: Назва», «ПДП Назва», перша назва в “лапках”
    zp = f'SEARCH("ЗПМ «",{B})'
    c1 = f'IFERROR(MID({B},{zp}+5,SEARCH("»",{B},{zp})-{zp}-5),"")'
    c2 = f'IFERROR(LEFT(MID({B},SEARCH(") ",{B})+2,SEARCH(" пілот",{B})-SEARCH(") ",{B})-2),25),"")'
    ek = f'MID({B},SEARCH("екіпаж: ",{B})+8,30)'
    c3 = f'IFERROR(LEFT({ek},SEARCH(" ",{ek}&" ")-1),"")'
    pd = f'MID({B},SEARCH("ПДП ",{B})+4,30)'
    c4 = f'IFERROR("ПДП "&LEFT({pd},SEARCH(" ",{pd}&" ")-1),"")'
    qo = f'SEARCH("“",{B})'
    qt = f'LEFT(MID({B},{qo}+1,SEARCH("”",{B},{qo})-{qo}-1),20)'
    c5 = f'IFERROR(LEFT({qt},SEARCH(CHAR(34),{qt}&CHAR(34))-1),"")'
    so = f'SEARCH(CHAR(34),{B})'
    c6 = f'IFERROR(LEFT(MID({B},{so}+1,SEARCH(CHAR(34),{B},{so}+1)-{so}-1),20),"")'
    crew = (f'IF({c1}<>"",{c1},IF({c2}<>"",{c2},IF({c3}<>"",{c3},IF({c4}<>"",{c4},'
            f'IF({c5}<>"",{c5},{c6})))))')
    for ch in ["«", "»", "“", "”"]:
        crew = f'SUBSTITUTE({crew},"{ch}","")'
    crew = f'TRIM(SUBSTITUTE(SUBSTITUTE({crew},CHAR(34),""),",",""))'

    def number_before_bort(word, span, unit=" борт"):
        w = f'MID({B},SEARCH("{word}",{B}),{span})'
        return f'VALUE(TRIM(RIGHT(SUBSTITUTE(LEFT({w},SEARCH("{unit}",{w})-1)," ",REPT(" ",20)),20)))'

    res = f"{JC['result']}{r}"
    guard = lambda body: f'=IF({B}="","",{body})'
    return {
        "time": guard(time_of(B)),
        "type": guard(kw_match("targets", seg, patrol)),
        "place": guard(f"IF({m}=0,{before_region},{after_marker})"),
        "grid": guard(grid),
        "means": guard(kw_match("means", B)),
        "ammo": guard(ammo),
        "crew": guard(crew),
        "result": guard(kw_match("results", B)),
        "spent": guard(f"IFERROR({number_before_bort('витрат', 40)},"
                       f"IFERROR({number_before_bort('витрат', 20, ' fpv')},0))"),
        "lost": guard(f'IFERROR({number_before_bort("втрати:", 25)},IF({res}="{LOST}",1,0))'),
        "time_of": time_of,
    }


def strike_formulas(r):
    R = f"{SC['text']}{r}"
    tail = f'MID({R},SEARCH(" по ",{R})+4,60)'
    end = "MIN(" + ",".join(f'IFERROR(SEARCH("{x}",{tail}),61)' for x in [" кв", " (", ". ", ","]) + ")"
    qty = (f'IFERROR(VALUE(TRIM(RIGHT(SUBSTITUTE(SUBSTITUTE(LEFT({R},SEARCH("од.",{R})-1),"("," "),'
           f'" ",REPT(" ",20)),20))),1)')
    guard = lambda body: f'=IF({R}="","",{body})'
    return {
        "time": guard(parse_formulas(r)["time_of"](R)),
        "kind": guard(kw_match("strike", R, str(FALLBACK_ROW))),
        "enemy": guard(kw_match("enemy", R)),
        "qty": guard(qty),
        "target": guard(f'IFERROR(TRIM(LEFT({tail},{end}-1)),"")'),
        "losses": guard(f'IF(ISNUMBER(SEARCH("без втрат",{R})),0,"")'),
    }


EXAMPLE_REPORT = ("07:24 екіпаж (FPV-перехоплювачів) «Позивний» пілот Прізвище, в р-ні н.п. Приклад, "
                  "виявлено БпЛА 'Zala' - 1 (37U XX 00000 00000). Застосовано FPV "
                  "дрон-перехоплювач мультироторного типу. Ціль знищено. Витрата: 1 борт, ЕД - 1 од.")
EXAMPLE_STRIKE = "12:19 удар Молнія (1 од.) по позиції «Приклад» (37U XX 00000 00000). Без втрат."


def build_unit(wb, unit, idx, radars):
    ws = wb.create_sheet(unit["code"])
    ws.sheet_properties.tabColor = C_TEAL
    page(ws)
    widths(ws, {"A": 10, "B": 60, "C": 8, "D": 16, "E": 16, "F": 19, "G": 18, "H": 16, "I": 16,
                "J": 22, "K": 10, "L": 10, "M": 22, "N": 8, "O": 13, "P": 2, "Q": 5, "R": 50,
                "S": 8, "T": 8, "U": 16, "V": 7, "W": 22, "X": 9})

    # --- шапка з шевроном
    band(ws, "A1")
    title(ws, "B1:X1", f'="ДОБОВА ДОПОВІДЬ ППО — "&{U_NAME}', size=16)
    ws.row_dimensions[1].height = 50
    emblem(ws, unit.get("emblem"), 0, 0, 70, 60, off_y=3)
    label(ws, "A2", "Підрозділ")
    inp(ws, U_NAME, unit["name"])
    ws[U_NAME].font = font(11, True, C_NAVY)
    label(ws, "C2:D2", "Дата звіту")
    auto(ws, U_DATE, f'=IF({q(SUM_SHEET)}!$C$3="","",{q(SUM_SHEET)}!$C$3)', fmt="dd.mm.yyyy", bold=True)
    ws["F2"] = "← береться з аркуша «Зведення»"
    style(ws, "F2:I2", fnt=font(8, italic=True, color="5B6B7F"), border=None, align=LEFT, merge=True)
    label(ws, "A3", "Черговий")
    inp(ws, U_DUTY)
    label(ws, "C3:D3", "Зміна з – по")
    inp(ws, U_FROM, fmt="hh:mm", align=CENTER)
    inp(ws, U_TO, fmt="hh:mm", align=CENTER)
    time_dv(ws, [U_FROM, U_TO])

    # --- підсумок (авто)
    jr = lambda key: f"${JC[key]}${J_FIRST}:${JC[key]}${J_LAST}"
    sr = lambda key: f"${SC[key]}${J_FIRST}:${SC[key]}${S_LAST}"
    section(ws, "A5:O5", "ПІДСУМОК ЗА ДОБУ — рахується автоматично з журналу")
    tile(ws, "A6:B6", "A7:B7", "Доповідей у журналі", f'=COUNTIF({jr("text")},"?*")')
    tile(ws, "C6:D6", "C7:D7", "Виявлено цілей", f'=COUNTIFS({jr("type")},"?*",{jr("type")},"<>{PATROL}")')
    tile(ws, "E6:F6", "E7:F7", "Застосувань", f'=COUNTIFS({jr("means")},"?*")')
    tile(ws, "G6:H6", "G7:H7", "Знищено", f'=COUNTIFS({jr("result")},"{DESTROYED}")')
    tile(ws, "I6:J6", "I7:J7", "Ефективність", '=IF(E7=0,"–",G7/E7)', fmt=PCT)
    tile(ws, "K6:L6", "K7:L7", "Витрачено бортів", f'=SUM({jr("spent")})')
    tile(ws, "M6:N6", "M7:N7", "Втрачено бортів", f'=SUM({jr("lost")})')
    tile(ws, "O6", "O7", "Перевірити ⚠", f'=COUNTIF({jr("check")},"⚠*")')
    section(ws, "Q5:X5", "УДАРИ ПРОТИВНИКА — підсумок", color=C_SLATE)
    tile(ws, "Q6:R6", "Q7:R7", "РУ", f'=COUNTIFS({sr("kind")},"РУ")')
    tile(ws, "S6:T6", "S7:T7", "АУ", f'=COUNTIFS({sr("kind")},"АУ")')
    tile(ws, "U6", "U7", "УДК", f'=COUNTIFS({sr("kind")},"УДК")')
    tile(ws, "V6:W6", "V7:W7", "Засобів противника, од.", f'=SUM({sr("qty")})')
    tile(ws, "X6", "X7", "Втрати о/с", f'=SUM({sr("losses")})')
    ws.row_dimensions[7].height = 32
    ws.conditional_formatting.add("O7", CellIsRule(operator="greaterThan", formula=["0"],
                                                    font=Font(name=FONT, size=18, bold=True, color="B45309")))

    # --- РЛС: наявність і робота
    r0, r1, t = RADAR_ROWS[0], RADAR_ROWS[-1], RADAR_TOTAL
    section(ws, "A9:F9", "СИЛИ І ЗАСОБИ — РЛС")
    section(ws, "H9:O9", "РОБОТА РЛС ЗА ДОБУ")
    for col, h in zip("ABCDEF", ["№", "Тип РЛС", "Всього", "Справні", "На ремонті", "На складі"]):
        header(ws, f"{col}10", h)
    for col, h in zip("HIJKL", ["Позивний РЛС", "Тип РЛС", "Виявлено цілей", "Включень", "Час роботи (год:хв)"]):
        header(ws, f"{col}10", h)
    header(ws, "M10:O10", "Періоди готовності №1 (з – по)")
    ws.row_dimensions[10].height = 28
    for i, r in enumerate(RADAR_ROWS):
        auto(ws, f"A{r}", i + 1)
        inp(ws, f"B{r}", radars[i] if i < len(radars) else None)
        for col in "CDEF":
            inp(ws, f"{col}{r}", fmt="0", align=CENTER)
        inp(ws, f"H{r}")
        inp(ws, f"I{r}")
        inp(ws, f"J{r}", fmt="0", align=CENTER)
        inp(ws, f"K{r}", fmt="0", align=CENTER)
        inp(ws, f"L{r}", fmt="[h]:mm", align=CENTER)
        inp(ws, f"M{r}:O{r}")
    label(ws, f"A{t}:B{t}", "Разом")
    for col in "CDEF":
        auto(ws, f"{col}{t}", f"=SUM({col}{r0}:{col}{r1})", fmt=NUM, bold=True)
    label(ws, f"H{t}:I{t}", "Разом")
    auto(ws, f"J{t}", f"=SUM(J{r0}:J{r1})", fmt=NUM, bold=True)
    auto(ws, f"K{t}", f"=SUM(K{r0}:K{r1})", fmt=NUM, bold=True)
    auto(ws, f"L{t}", f"=SUM(L{r0}:L{r1})", fmt="[h]:mm", bold=True)
    auto(ws, f"M{t}:O{t}")
    list_dv(ws, "lst_radar", [f"B{r0}:B{r1}", f"I{r0}:I{r1}"])
    num_dv(ws, [f"C{r0}:F{r1}", f"J{r0}:K{r1}"])

    # --- екіпажі
    c0, c1, t = CREW_ROWS[0], CREW_ROWS[-1], CREW_TOTAL
    section(ws, "A17:O17", "ЕКІПАЖІ ПЕРЕХОПЛЮВАЧІВ — НАЯВНІСТЬ")
    header(ws, "A18", "№")
    header(ws, "B18", "Екіпаж (позивний)")
    header(ws, "C18:D18", "Тип перехоплювача")
    for col, h in zip("EFGHIJ", ["Бортів на позиції (день)", "Бортів на позиції (ніч)", "Бортів на складі",
                                 "Швидкість, км/год", "Стан", "Район (Н.П.)"]):
        header(ws, f"{col}18", h)
    header(ws, "K18:L18", "Дата відновлення")
    header(ws, "M18:O18", "Примітка")
    ws.row_dimensions[18].height = 28
    for i, r in enumerate(CREW_ROWS):
        auto(ws, f"A{r}", i + 1)
        inp(ws, f"B{r}")
        inp(ws, f"C{r}:D{r}")
        for col in "EFG":
            inp(ws, f"{col}{r}", fmt="0", align=CENTER)
        inp(ws, f"H{r}", align=CENTER)
        inp(ws, f"I{r}")
        inp(ws, f"J{r}")
        inp(ws, f"K{r}:L{r}", fmt="dd.mm.yyyy", align=CENTER)
        inp(ws, f"M{r}:O{r}")
    label(ws, f"A{t}", "Разом")
    auto(ws, f"B{t}", f"=COUNTA(B{c0}:B{c1})", fmt='0" екіпаж."', bold=True)
    auto(ws, f"C{t}:D{t}")
    for col in "EFG":
        auto(ws, f"{col}{t}", f"=SUM({col}{c0}:{col}{c1})", fmt=NUM, bold=True)
    auto(ws, f"H{t}")
    auto(ws, f"I{t}", f'=COUNTIFS(I{c0}:I{c1},"{READY}")', fmt='0" боєгот."', bold=True)
    auto(ws, f"J{t}")
    auto(ws, f"K{t}:L{t}")
    auto(ws, f"M{t}:O{t}")
    list_dv(ws, "lst_status", [f"I{c0}:I{c1}"])
    num_dv(ws, [f"E{c0}:G{c1}"])
    date_dv(ws, [f"K{c0}:K{c1}"])

    # --- журнал цілей і удари противника
    section(ws, f"A{J_HEAD - 1}:O{J_HEAD - 1}",
            "ЖУРНАЛ ЦІЛЕЙ — вставте текст доповіді у стовпець B, решта заповниться сама "
            "(блакитні поля можна виправити)")
    section(ws, f"Q{J_HEAD - 1}:X{J_HEAD - 1}", "УДАРИ ПРОТИВНИКА (РУ / АУ / УДК) — вставте текст у стовпець R",
            color=C_SLATE)
    jcols = {"num": "№", "text": "Текст доповіді (вставте сюди)", "time": "Час", "type": "Тип цілі",
             "place": "Н.П. / район", "grid": "Квадрат (MGRS)", "means": "Засіб ураження",
             "ammo": "Боєприпас", "crew": "Екіпаж / позиція", "result": "Результат",
             "spent": "Витрачено бортів", "lost": "Втрачено бортів", "note": "Примітка",
             "group": "Група засобу", "check": "Перевірка"}
    scols = {"num": "№", "text": "Текст доповіді про удар (вставте сюди)", "time": "Час",
             "kind": "Тип удару", "enemy": "Засіб противника", "qty": "К-сть, од.",
             "target": "Об'єкт / район", "losses": "Втрати о/с"}
    for key, h in jcols.items():
        header(ws, f"{JC[key]}{J_HEAD}", h)
    for key, h in scols.items():
        header(ws, f"{SC[key]}{J_HEAD}", h)
    ws.row_dimensions[J_HEAD].height = 30

    group = (lambda r: f'=IF({JC["means"]}{r}="","",IFERROR(INDEX({rl("means", True)},'
                       f'MATCH({JC["means"]}{r},{rl("means")},0)),"{G_OTHER}"))')
    TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
    TOPC = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for r in range(J_EXAMPLE, J_LAST + 1):
        ex = r == J_EXAMPLE
        f = parse_formulas(r)
        J = {k: f"{v}{r}" for k, v in JC.items()}
        auto(ws, J["num"], "Приклад" if ex else
             f'=IF(AND({J["text"]}="",{J["type"]}=""),"",ROW()-{J_FIRST - 1})', align=TOPC)
        inp(ws, J["text"], EXAMPLE_REPORT if ex else None, align=TOP)
        parsed(ws, J["time"], f["time"], fmt="hh:mm", align=TOPC)
        for key in ["type", "place", "grid", "means", "ammo", "crew", "result"]:
            parsed(ws, J[key], f[key], align=TOP)
        parsed(ws, J["spent"], f["spent"], fmt="0", align=TOPC)
        parsed(ws, J["lost"], f["lost"], fmt="0", align=TOPC)
        inp(ws, J["note"], align=TOP)
        auto(ws, J["group"], group(r), align=TOPC)
        auto(ws, J["check"], f'=IF({J["text"]}="","",IF(OR({J["type"]}="",{J["result"]}=""),'
                             f'"⚠ перевірте","✓"))', align=TOPC)
        if r <= S_LAST:
            g = strike_formulas(r)
            S = {k: f"{v}{r}" for k, v in SC.items()}
            auto(ws, S["num"], "Приклад" if ex else
                 f'=IF(AND({S["text"]}="",{S["kind"]}=""),"",ROW()-{J_FIRST - 1})', align=TOPC)
            inp(ws, S["text"], EXAMPLE_STRIKE if ex else None, align=TOP)
            parsed(ws, S["time"], g["time"], fmt="hh:mm", align=TOPC)
            parsed(ws, S["kind"], g["kind"], align=TOPC)
            parsed(ws, S["enemy"], g["enemy"], align=TOP)
            parsed(ws, S["qty"], g["qty"], fmt="0", align=TOPC)
            parsed(ws, S["target"], g["target"], align=TOP)
            parsed(ws, S["losses"], g["losses"], fmt="0", align=TOPC)

    # рядок-приклад (не враховується у підсумках)
    e = J_EXAMPLE
    for c in [*JC.values(), *SC.values()]:
        ws[f"{c}{e}"].font = font(9, italic=True, color="7F7F7F")
        ws[f"{c}{e}"].fill = fill(C_EXAMPLE)
    ws[f"A{e}"].comment = Comment("Рядок-приклад: не враховується у підсумках. Можна очистити.", "ППО")

    ws.add_table(Table(displayName=f"tblJournal{idx}", ref=f"A{J_HEAD}:O{J_LAST}",
                       tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=False)))
    ws.add_table(Table(displayName=f"tblStrikes{idx}", ref=f"Q{J_HEAD}:X{S_LAST}",
                       tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=False)))

    rng = lambda key, last=J_LAST, cols=JC: f"{cols[key]}{J_EXAMPLE}:{cols[key]}{last}"
    list_dv(ws, "lst_targets", [rng("type")])
    list_dv(ws, "lst_means", [rng("means")])
    # боєприпаси — лише для обраного засобу (залежний список)
    g0 = f"${JC['means']}{J_EXAMPLE}"
    AW, AA = rl("ammo"), rl("ammo", True)
    first = f"{q(REF)}!${RL['ammo'][1]}$5"
    dv = DataValidation(type="list", allow_blank=True, errorStyle="warning",
                        formula1=f'IF(OR({g0}="",COUNTIF({AW},{g0})=0),{AA},'
                                 f'OFFSET({first},MATCH({g0},{AW},0)-1,0,COUNTIF({AW},{g0}),1))')
    dv.error, dv.errorTitle = "Такого боєприпасу немає в довіднику для цього засобу.", "Боєприпас"
    ws.add_data_validation(dv)
    dv.add(rng("ammo"))
    list_dv(ws, "lst_results", [rng("result")])
    num_dv(ws, [rng("spent"), rng("lost")], hi=99)
    time_dv(ws, [rng("time"), rng("time", S_LAST, SC)])
    list_dv(ws, "lst_strike", [rng("kind", S_LAST, SC)],
            prompt="РУ — ракетний, АУ — авіаційний, УДК — дрон-камікадзе")
    list_dv(ws, "lst_enemy", [rng("enemy", S_LAST, SC)])
    num_dv(ws, [rng("qty", S_LAST, SC), rng("losses", S_LAST, SC)])

    # підсвічування: результат, нерозпізнані поля, позначка перевірки
    res = f"{JC['result']}{J_FIRST}:{JC['result']}{J_LAST}"
    for val, bg, fg in [(DESTROYED, "D8F0DE", "1E6B34"), ("Пошкоджено", "FFF1C2", "7A5A00"),
                        ("Не знищено", "FAD9DC", "8E1B26"), (LOST, "FAD9DC", "8E1B26")]:
        ws.conditional_formatting.add(res, CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                      fill=fill(bg), font=Font(name=FONT, color=fg, bold=True)))
    for key in ["type", "result"]:
        col = JC[key]
        ws.conditional_formatting.add(
            f"{col}{J_FIRST}:{col}{J_LAST}",
            FormulaRule(formula=[f'AND(${JC["text"]}{J_FIRST}<>"",{col}{J_FIRST}="")'], fill=fill("FDE2C4")))
    chk = f"{JC['check']}{J_FIRST}:{JC['check']}{J_LAST}"
    ws.conditional_formatting.add(chk, CellIsRule(operator="equal", formula=['"⚠ перевірте"'],
                                                  fill=fill("FDE2C4"), font=Font(name=FONT, color="9A3412", bold=True)))
    ws.conditional_formatting.add(chk, CellIsRule(operator="equal", formula=['"✓"'],
                                                  font=Font(name=FONT, color="1E6B34", bold=True)))
    protect(ws)
    return ws


# ================================================================ Зведення
SUM_SHEET = "Зведення"
U_FIRST = 11                     # перший рядок підрозділу у таблиці «Бойова робота»
U_LAST = U_FIRST + UNIT_SLOTS - 1


def ind(code, rng):
    return f'INDIRECT("\'"&{code}&"\'!{rng}")'


def jcol(key):
    return f"${JC[key]}${J_FIRST}:${JC[key]}${J_LAST}"


def scol(key):
    return f"${SC[key]}${J_FIRST}:${SC[key]}${S_LAST}"


def build_summary(wb, cfg):
    ws = wb.create_sheet(SUM_SHEET)
    ws.sheet_properties.tabColor = C_NAVY
    page(ws)
    # стовпці таблиці «Бойова робота»: E виявлено, F застосовано, по 2 на групу засобів, далі підсумки
    gcols = [(get_column_letter(7 + 2 * i), get_column_letter(8 + 2 * i)) for i in range(len(GROUPS))]
    nxt = 7 + 2 * len(GROUPS)
    C_DES, C_EFF, C_SPENT, C_LOST, C_DUTY = (get_column_letter(nxt + k) for k in range(5))
    last = C_DUTY
    widths(ws, {"A": 5, "B": 11, "C": 28, "D": 13, **{get_column_letter(c): 10 for c in range(5, nxt + 4)},
                C_DUTY: 30})
    band(ws, "A1:B1")
    title(ws, f"C1:{last}1", f'="ЗВЕДЕННЯ ППО — "&{q(REF)}!$B$2', size=16)
    ws.row_dimensions[1].height = 50
    emblem(ws, cfg.get("org_emblem"), 0, 0, 120, 60, off_y=3)
    label(ws, "A3:B3", "Дата звіту")
    inp(ws, "C3", fmt="dd.mm.yyyy", align=CENTER)
    ws["C3"].font = font(12, True, C_NAVY)
    date_dv(ws, ["C3"])
    ws["A4"] = "← вкажіть дату; усе інше рахується автоматично з аркушів підрозділів"
    style(ws, "A4:D4", fnt=font(8, italic=True, color="5B6B7F"), border=None, align=LEFT, merge=True)

    tot = U_LAST + 1
    strike_tot = 45
    tile(ws, "E3:F3", "E4:F5", "Виявлено цілей", f"=E{tot}")
    tile(ws, "G3:H3", "G4:H5", "Застосувань", f"=F{tot}")
    tile(ws, "I3:J3", "I4:J5", "Знищено", f"={C_DES}{tot}")
    tile(ws, "K3:L3", "K4:L5", "Ефективність", f"={C_EFF}{tot}", fmt=PCT)
    tile(ws, "M3:N3", "M4:N5", "Витрачено бортів", f"={C_SPENT}{tot}")
    tile(ws, "O3:P3", "O4:P5", "Втрачено бортів", f"={C_LOST}{tot}")
    tile(ws, "Q3:R3", "Q4:R5", "Ударів противника", f"=SUM(D{strike_tot}:F{strike_tot})")
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 20

    def unit_cols(r, i):
        """№, код, назва — спільні для трьох таблиць."""
        n = 5 + i
        auto(ws, f"A{r}", f'=IF(B{r}="","",{i + 1})')
        auto(ws, f"B{r}", f'=IF({q(REF)}!$A${n}="","",{q(REF)}!$A${n})', bold=True)
        auto(ws, f"C{r}", f'=IF(B{r}="","",IF({q(REF)}!$B${n}="",B{r},{q(REF)}!$B${n}))', align=LEFT)

    # --- 1. Бойова робота
    section(ws, f"A8:{last}8", "БОЙОВА РОБОТА ПО ПІДРОЗДІЛАХ")
    heads = [("A9:A10", "№"), ("B9:B10", "Код"), ("C9:C10", "Підрозділ"), ("D9:D10", "Аркуш"),
             ("E9:E10", "Виявлено цілей"), ("F9:F10", "Застосовано всього"),
             (f"{C_DES}9:{C_DES}10", "Знищено всього"), (f"{C_EFF}9:{C_EFF}10", "Ефектив-ність"),
             (f"{C_SPENT}9:{C_SPENT}10", "Витрачено бортів"), (f"{C_LOST}9:{C_LOST}10", "Втрачено бортів"),
             (f"{C_DUTY}9:{C_DUTY}10", "Черговий")]
    for (ca, cd), (_, glabel) in zip(gcols, GROUPS):
        heads.append((f"{ca}9:{cd}9", glabel))
    for rng, t in heads:
        header(ws, rng, t)
    for ca, cd in gcols:
        header(ws, f"{ca}10", "застос.")
        header(ws, f"{cd}10", "знищ.")
    ws.row_dimensions[9].height = 30

    for i in range(UNIT_SLOTS):
        r = U_FIRST + i
        unit_cols(r, i)
        code = f"$B{r}"
        auto(ws, f"D{r}", f'=IF({code}="","",IF(ISREF({ind(code, "A1")}),"✓","немає аркуша"))')
        ok = f'$D{r}<>"✓"'
        T, G, R = ind(code, jcol("type")), ind(code, jcol("group")), ind(code, jcol("result"))
        f = {
            "E": f'COUNTIFS({T},"?*",{T},"<>{PATROL}")',
            "F": f'COUNTIFS({ind(code, jcol("means"))},"?*")',
            C_DES: f'COUNTIFS({R},"{DESTROYED}")',
            C_SPENT: f'SUM({ind(code, jcol("spent"))})',
            C_LOST: f'SUM({ind(code, jcol("lost"))})',
        }
        for (ca, cd), (g, _) in zip(gcols, GROUPS):
            f[ca] = f'COUNTIFS({G},"{g}")'
            f[cd] = f'COUNTIFS({G},"{g}",{R},"{DESTROYED}")'
        for col, expr in f.items():
            auto(ws, f"{col}{r}", f'=IF({ok},"",{expr})', fmt=NUM, bold=col == C_DES)
        auto(ws, f"{C_EFF}{r}", f'=IF({ok},"",IF(N(F{r})=0,"–",{C_DES}{r}/F{r}))', fmt=PCT)
        auto(ws, f"{C_DUTY}{r}", f'=IF({ok},"",{ind(code, U_DUTY)}&"")', align=LEFT)
    label(ws, f"A{tot}:D{tot}", "РАЗОМ")
    for c in range(5, nxt + 4):
        col = get_column_letter(c)
        if col != C_EFF:
            auto(ws, f"{col}{tot}", f"=SUM({col}{U_FIRST}:{col}{U_LAST})", fmt=NUM, bold=True)
    auto(ws, f"{C_EFF}{tot}", f'=IF(F{tot}=0,"–",{C_DES}{tot}/F{tot})', fmt=PCT, bold=True)
    auto(ws, f"{C_DUTY}{tot}")
    style(ws, f"A{tot}:{last}{tot}", f=fill("DDE3EA"))
    ws.conditional_formatting.add(f"{C_DES}{U_FIRST}:{C_DES}{U_LAST}",
                                  DataBarRule(start_type="num", start_value=0, end_type="max",
                                              color="7FB3A8", showValue=True))
    ws.conditional_formatting.add(f"D{U_FIRST}:D{U_LAST}",
                                  CellIsRule(operator="equal", formula=['"немає аркуша"'],
                                             fill=fill("FAD9DC"), font=Font(name=FONT, color="8E1B26")))

    # --- 2. Удари противника
    s_head = 28
    s_first = s_head + 2
    assert s_first + UNIT_SLOTS == strike_tot
    section(ws, f"A{s_head}:{last}{s_head}", "УДАРИ ПРОТИВНИКА ПО ПІДРОЗДІЛАХ  (РУ / АУ / УДК — кількість "
                                              "ударів; засоби противника — кількість одиниць)", color=C_SLATE)
    h = s_head + 1
    c_loss = get_column_letter(7 + ENEMY_SLOTS)          # після стовпців засобів противника
    for col, t in [("A", "№"), ("B", "Код"), ("C", "Підрозділ"), ("D", "РУ"), ("E", "АУ"), ("F", "УДК"),
                   (c_loss, "Втрати о/с")]:
        header(ws, f"{col}{h}", t)
    for k in range(ENEMY_SLOTS):
        col = get_column_letter(7 + k)
        en = RL["enemy"][0]
        header(ws, f"{col}{h}", f'=IF({q(REF)}!${en}${5 + k}="","",{q(REF)}!${en}${5 + k})')
    ws.row_dimensions[h].height = 30
    for i in range(UNIT_SLOTS):
        r = s_first + i
        unit_cols(r, i)
        code, ok = f"$B{r}", f'$D${U_FIRST + i}<>"✓"'
        for col, kind in zip("DEF", ["РУ", "АУ", "УДК"]):
            auto(ws, f"{col}{r}", f'=IF({ok},"",COUNTIFS({ind(code, scol("kind"))},"{kind}"))', fmt=NUM)
        for k in range(ENEMY_SLOTS):
            col = get_column_letter(7 + k)
            auto(ws, f"{col}{r}", f'=IF(OR({ok},{col}${h}=""),"",'
                                  f'SUMIFS({ind(code, scol("qty"))},{ind(code, scol("enemy"))},{col}${h}))', fmt=NUM)
        auto(ws, f"{c_loss}{r}", f'=IF({ok},"",SUM({ind(code, scol("losses"))}))', fmt=NUM)
    label(ws, f"A{strike_tot}:C{strike_tot}", "РАЗОМ")
    for c in range(4, 8 + ENEMY_SLOTS):
        col = get_column_letter(c)
        auto(ws, f"{col}{strike_tot}", f"=SUM({col}{s_first}:{col}{strike_tot - 1})", fmt=NUM, bold=True)
    style(ws, f"A{strike_tot}:{c_loss}{strike_tot}", f=fill("DDE3EA"))

    # --- 3. Сили і засоби
    f_head = strike_tot + 2
    h = f_head + 1
    f_first = f_head + 3
    section(ws, f"A{f_head}:{last}{f_head}", "СИЛИ І ЗАСОБИ ПО ПІДРОЗДІЛАХ")
    for rng, t in [(f"A{h}:A{h + 1}", "№"), (f"B{h}:B{h + 1}", "Код"), (f"C{h}:C{h + 1}", "Підрозділ"),
                   (f"D{h}:D{h + 1}", "Зміна"), (f"E{h}:H{h}", "РЛС (наявність)"),
                   (f"I{h}:K{h}", "Робота РЛС"), (f"L{h}:M{h}", "Екіпажі"),
                   (f"N{h}:P{h}", "Борти перехоплювачів"), (f"Q{h}:{last}{h + 1}", "Черговий")]:
        header(ws, rng, t)
    for col, t in zip("EFGHIJKLMNOP", ["всього", "справні", "ремонт", "склад", "виявлено цілей",
                                        "включень", "год роботи", "всього", "боєготові", "позиція день",
                                        "позиція ніч", "склад"]):
        header(ws, f"{col}{h + 1}", t)
    ws.row_dimensions[h + 1].height = 26
    rt, ct = RADAR_TOTAL, CREW_TOTAL
    src = {"E": f"C{rt}", "F": f"D{rt}", "G": f"E{rt}", "H": f"F{rt}", "I": f"J{rt}", "J": f"K{rt}",
           "K": f"L{rt}", "L": f"B{ct}", "M": f"I{ct}", "N": f"E{ct}", "O": f"F{ct}", "P": f"G{ct}"}
    f_tot = f_first + UNIT_SLOTS
    for i in range(UNIT_SLOTS):
        r = f_first + i
        unit_cols(r, i)
        code, ok = f"$B{r}", f'$D${U_FIRST + i}<>"✓"'
        a, b = ind(code, U_FROM), ind(code, U_TO)
        auto(ws, f"D{r}", f'=IF({ok},"",IF(AND({a}&""="",{b}&""=""),"",'
                          f'IFERROR({hhmm(a)}&"–"&{hhmm(b)},"")))')
        for col, cell in src.items():
            auto(ws, f"{col}{r}", f'=IF({ok},"",N({ind(code, cell)}))',
                 fmt="[h]:mm" if col == "K" else NUM)
        auto(ws, f"Q{r}:{last}{r}", f'=IF({ok},"",{ind(code, U_DUTY)}&"")', align=LEFT)
    label(ws, f"A{f_tot}:D{f_tot}", "РАЗОМ")
    for col in "EFGHIJKLMNOP":
        auto(ws, f"{col}{f_tot}", f"=SUM({col}{f_first}:{col}{f_tot - 1})",
             fmt="[h]:mm" if col == "K" else NUM, bold=True)
    auto(ws, f"Q{f_tot}:{last}{f_tot}")
    style(ws, f"A{f_tot}:{last}{f_tot}", f=fill("DDE3EA"))

    ws.freeze_panes = "A8"
    protect(ws)
    return ws


# ================================================================ Розрахунок + За типами цілей
CALC = "Розрахунок"
CALC_BLOCKS = ["Виявлено", "Застосовано", *[f"{g} знищено" for g, _ in GROUPS], "Знищено всього"]
BLOCK_W = UNIT_SLOTS + 2         # підрозділи + «Разом» + проміжок
T_FIRST = 5
T_LAST = T_FIRST + TYPE_SLOTS - 1


def calc_col(block, unit=None):
    """Стовпець у «Розрахунку»: unit=None → стовпець «Разом» блоку."""
    start = 2 + block * BLOCK_W
    return get_column_letter(start + (UNIT_SLOTS if unit is None else unit))


def build_calc(wb):
    ws = wb.create_sheet(CALC)
    ws.sheet_properties.tabColor = "8C96A3"
    ws.sheet_view.showGridLines = False
    title(ws, f"A1:{calc_col(len(CALC_BLOCKS) - 1)}1",
          "РОЗРАХУНОК — службовий аркуш (рахується автоматично, не редагується)", size=11)
    ws.column_dimensions["A"].width = 26
    header(ws, "A4", "Тип цілі")
    extra = [lambda u: "", lambda u: f',{ind(u, jcol("means"))},"?*"']
    extra += [(lambda g: lambda u: f',{ind(u, jcol("group"))},"{g}",{ind(u, jcol("result"))},"{DESTROYED}"')(g)
              for g, _ in GROUPS]
    extra += [lambda u: f',{ind(u, jcol("result"))},"{DESTROYED}"']
    for t in range(TYPE_SLOTS):
        r = T_FIRST + t
        auto(ws, f"A{r}", f'=IF({q(REF)}!$D${5 + t}="","",{q(REF)}!$D${5 + t})', align=LEFT)
    for b, name in enumerate(CALC_BLOCKS):
        section(ws, f"{calc_col(b, 0)}3:{calc_col(b)}3", name, color=C_SLATE)
        for u in range(UNIT_SLOTS):
            col = calc_col(b, u)
            ws.column_dimensions[col].width = 7
            header(ws, f"{col}4", f'=IF({q(REF)}!$A${5 + u}="","",{q(REF)}!$A${5 + u})')
            code = f"{col}$4"
            for t in range(TYPE_SLOTS):
                r = T_FIRST + t
                auto(ws, f"{col}{r}", f'=IF(OR($A{r}="",{code}=""),0,'
                                      f'IFERROR(COUNTIFS({ind(code, jcol("type"))},$A{r}{extra[b](code)}),0))',
                     fmt=NUM)
        tc = calc_col(b)
        ws.column_dimensions[tc].width = 8
        header(ws, f"{tc}4", "Разом")
        for t in range(TYPE_SLOTS):
            r = T_FIRST + t
            auto(ws, f"{tc}{r}", f"=SUM({calc_col(b, 0)}{r}:{calc_col(b, UNIT_SLOTS - 1)}{r})",
                 fmt=NUM, bold=True)
        ws.column_dimensions[get_column_letter(2 + b * BLOCK_W + UNIT_SLOTS + 1)].width = 2
    ws.freeze_panes = "B5"
    protect(ws)
    return ws


# ================================================================ Доповідь
REPORT = "Доповідь"
ALL_UNITS = "Усі підрозділи"
# Показники: ключ, назва, джерело (S — удари противника, D — знищено, A — застосування), що рахувати
INDICATORS = [
    ("ru", "Ракетні удари (к-сть)", "S", "РУ"),
    ("au", "Авіаційні удари (к-сть)", "S", "АУ"),
    ("kab", "КАБ, од.", "S", "КАБ"),
    ("kar", "КАР, од.", "S", "КАР"),
    ("h_shahed", "Влучання: шахеди", "S", "Шахед"),
    ("h_gerbera", "Влучання: гербери", "S", "Гербера"),
    ("h_geran", "Влучання: герань", "S", "Герань"),
    ("h_banderol", "Влучання: бандероль", "S", "Бандероль"),
    ("h_italmas", "Влучання: італмас", "S", "Італмас"),
    ("h_lancet", "Влучання: ланцет", "S", "Ланцет"),
    ("h_molniya", "Влучання: молнія", "S", "Молнія"),
    ("h_fpv", "Влучання: FPV", "S", "FPV"),
    ("d_shahed", "Знищено: шахеди", "D", "Shahed"),
    ("d_gerbera", "Знищено: гербери", "D", "Гербера"),
    ("d_italmas", "Знищено: італмас", "D", "Італмас"),
    ("d_sokil", "Знищено: сокіл", "D", "Сокіл"),
    ("d_v2u", "Знищено: V2U", "D", "V2U"),
    ("d_lancet", "Знищено: ланцет", "D", "Ланцет"),
    ("d_molniya", "Знищено: молнія", "D", "Молнія"),
    ("d_fpv", "Знищено: FPV", "D", "FPV; FPV (оптоволокно)"),
    ("d_bomber", "Знищено: бомбер", "D", "Бомбер"),
    ("d_copter", "Знищено: коптерного типу", "D", "Mavic / Autel; DJI Matrice"),
    ("d_hexa", "Знищено: гексокоптер", "D", "Гексокоптер"),
    ("d_orlan", "Знищено: Орлан", "D", "Орлан-10/30"),
    ("d_zala", "Знищено: Зала", "D", "Zala"),
    ("d_skat", "Знищено: Скат", "D", "Скат"),
    ("d_supercam", "Знищено: Суперкам", "D", "Supercam"),
    ("d_kvo", "Знищено: КВО", "D", "КВО (Князь Віщий Олег)"),
    ("app", "Застосувань", "A", ""),
]
# Рядки доповіді: шаблон і значення {1}…{5}; «ключ.M» — всього, «ключ.N» — з них FPV-перехоплювачами
REPORT_LINES = [
    ("{1}", ["NAME"]),
    ("1) Ракетні удари - {1}", ["ru.M"]),
    ("2) авіаційні удари – {1} (каб-{2}, кар-{3})", ["au.M", "kab.M", "kar.M"]),
    ("3) дальні ударні БпЛА:", []),
    ("влучання – {1} (шахеди - {2}/гербери-{3}/герань-{4}/бандероль - {5})",
     ["h_shahed.M+h_gerbera.M+h_geran.M+h_banderol.M", "h_shahed.M", "h_gerbera.M", "h_geran.M", "h_banderol.M"]),
    ("знищено – {1} (шахеди – {2}/{3}, гербери – {4}/{5})",
     ["d_shahed.M+d_gerbera.M", "d_shahed.M", "d_shahed.N", "d_gerbera.M", "d_gerbera.N"]),
    ("4) ударні БпЛА:", []),
    ("влучання:", []),
    ("італмас – {1}", ["h_italmas.M"]),
    ("ланцет – {1}", ["h_lancet.M"]),
    ("молнія – {1}", ["h_molniya.M"]),
    ("FPV – {1}", ["h_fpv.M"]),
    ("знищення:", []),
    ("Італмас - {1}/{2}", ["d_italmas.M", "d_italmas.N"]),
    ("Сокіл - {1}/{2}", ["d_sokil.M", "d_sokil.N"]),
    ("V2U - {1}/{2}", ["d_v2u.M", "d_v2u.N"]),
    ("ланцет – {1}/{2}", ["d_lancet.M", "d_lancet.N"]),
    ("молнія – {1}/{2}", ["d_molniya.M", "d_molniya.N"]),
    ("FPV – {1}/{2}", ["d_fpv.M", "d_fpv.N"]),
    ("бомбер - {1}/{2}", ["d_bomber.M", "d_bomber.N"]),
    ("коптерного типу – {1}/{2}", ["d_copter.M", "d_copter.N"]),
    ("Гексокоптер - {1}/{2}", ["d_hexa.M", "d_hexa.N"]),
    ("5) розвідувальні БпЛА:", []),
    ("Знищено:", []),
    ("Орлан – {1}/{2}", ["d_orlan.M", "d_orlan.N"]),
    ("Зала – {1}/{2}", ["d_zala.M", "d_zala.N"]),
    ("Скат - {1}/{2}", ["d_skat.M", "d_skat.N"]),
    ("Суперкам – {1}/{2}", ["d_supercam.M", "d_supercam.N"]),
    ("КВО – {1}/{2}", ["d_kvo.M", "d_kvo.N"]),
    ("6) Застосувань – {1}/{2}", ["app.M", "app.N"]),
]
R_FIRST = 8
R_LINES = 40


def build_report(wb):
    import re

    ws = wb.create_sheet(REPORT)
    ws.sheet_properties.tabColor = C_TEAL
    page(ws)
    widths(ws, {"A": 2, "B": 66, "C": 2, "D": 56, "E": 7, "F": 7, "G": 7, "H": 7, "I": 7, "J": 2,
                "K": 28, "L": 34, "M": 9, "N": 9, "O": 2, "P": 16})
    title(ws, "A1:N1", "ФОРМУВАННЯ ДОПОВІДІ — текст складається автоматично з журналів підрозділів", size=14)
    ws.row_dimensions[1].height = 34
    label(ws, "B3", "Підрозділ (оберіть зі списку)")
    inp(ws, "D3", ALL_UNITS)
    ws["D3"].font = font(12, True, C_NAVY)
    label(ws, "B4", "Дата звіту")
    auto(ws, "D4", f'=IF({q(SUM_SHEET)}!$C$3="","",{q(SUM_SHEET)}!$C$3)', fmt="dd.mm.yyyy", bold=True, align=LEFT)
    ws["B5"] = (f"Виділіть клітинки B{R_FIRST}:B{R_FIRST + R_LINES - 1} → Ctrl+C → вставте в месенджер. "
                "Числа «x/y»: всього / з них FPV-перехоплювачами.")
    style(ws, "B5:N5", fnt=font(9, italic=True, color="5B6B7F"), border=None, align=LEFT, merge=True)

    # службові: список для вибору підрозділу і його номер у зведенні
    ws["P4"] = "службове"
    ws["P5"] = ALL_UNITS
    for i in range(UNIT_SLOTS):
        ws[f"P{6 + i}"] = f'=IF({q(REF)}!$A${5 + i}="","",{q(REF)}!$A${5 + i})'
    ws["P3"] = (f'=IF($D$3="{ALL_UNITS}",{UNIT_SLOTS + 1},'
                f'IFERROR(MATCH($D$3,{q(CALC)}!${calc_col(0, 0)}$4:${calc_col(0, UNIT_SLOTS - 1)}$4,0),0))')
    ws.column_dimensions["P"].hidden = True
    dv = DataValidation(type="list", formula1=f"$P$5:$P${5 + UNIT_SLOTS}", allow_blank=False)
    dv.error, dv.errorTitle = "Оберіть «Усі підрозділи» або код підрозділу.", "Підрозділ"
    ws.add_data_validation(dv)
    dv.add("D3")
    idx = "$P$3"

    header(ws, f"B{R_FIRST - 1}", "ТЕКСТ ДОПОВІДІ (копіюйте цей стовпець)")
    header(ws, f"D{R_FIRST - 1}", "Шаблон рядка — можна редагувати; {1}…{5} — числа праворуч")
    for k in range(5):
        header(ws, f"{get_column_letter(5 + k)}{R_FIRST - 1}", "{" + str(k + 1) + "}")
    for col, t in zip("KLMN", ["Показник", "Що рахувати (назви з «Довідників» через ;)", "Всього", "з них FPV"]):
        header(ws, f"{col}{R_FIRST - 1}", t)
    ws.row_dimensions[R_FIRST - 1].height = 30

    # --- показники
    RA = f"{q(CALC)}!$A${T_FIRST}:$A${T_LAST}"
    blk = lambda b: f"{q(CALC)}!${calc_col(b, 0)}${T_FIRST}:${calc_col(b)}${T_LAST}"
    b_all, b_fpv = len(CALC_BLOCKS) - 1, CALC_BLOCKS.index(f"{G_FPV} знищено")
    s_first, s_tot = 30, 45                       # таблиця ударів у «Зведенні»
    HD = f"{q(SUM_SHEET)}!$D${s_first - 1}:${get_column_letter(6 + ENEMY_SLOTS)}${s_first - 1}"
    ST = f"{q(SUM_SHEET)}!$D${s_first}:${get_column_letter(6 + ENEMY_SLOTS)}${s_tot}"
    pos = {}
    for i, (key, name, src, names) in enumerate(INDICATORS):
        r = R_FIRST + i
        pos[key] = r
        auto(ws, f"K{r}", name, align=LEFT)
        L = f'";"&SUBSTITUTE($L{r},"; ",";")&";"'
        if src == "A":
            auto(ws, f"L{r}", "застосування з «Зведення»", align=LEFT)
            m = f"=IF({idx}=0,0,N(INDEX({q(SUM_SHEET)}!$F${U_FIRST}:$F${U_LAST + 1},{idx})))"
            n = f"=IF({idx}=0,0,N(INDEX({q(SUM_SHEET)}!$G${U_FIRST}:$G${U_LAST + 1},{idx})))"
        elif src == "S":
            inp(ws, f"L{r}", names)
            m = f'=IF({idx}=0,0,SUMPRODUCT(ISNUMBER(SEARCH(";"&{HD}&";",{L}))*({HD}<>""),INDEX({ST},{idx},0)))'
            n = None
        else:
            inp(ws, f"L{r}", names)
            hit = f'ISNUMBER(SEARCH(";"&{RA}&";",{L}))*({RA}<>"")'
            m = f"=IF({idx}=0,0,SUMPRODUCT({hit},INDEX({blk(b_all)},0,{idx})))"
            n = f"=IF({idx}=0,0,SUMPRODUCT({hit},INDEX({blk(b_fpv)},0,{idx})))"
        auto(ws, f"M{r}", m, fmt="0", bold=True)
        auto(ws, f"N{r}", n, fmt="0")

    # --- рядки доповіді
    name = (f'=IF($D$3="{ALL_UNITS}",{q(REF)}!$B$2,IFERROR(INDEX({rl("units", True)},'
            f'MATCH($D$3,{rl("units")},0)),$D$3))')

    def value(expr):
        if expr == "NAME":
            return name
        return "=" + re.sub(r"(\w+)\.([MN])", lambda mm: f"${mm.group(2)}${pos[mm.group(1)]}", expr)

    for i in range(R_LINES):
        r = R_FIRST + i
        tpl, vals = REPORT_LINES[i] if i < len(REPORT_LINES) else (None, [])
        inp(ws, f"D{r}", tpl)
        for k in range(5):
            c = f"{get_column_letter(5 + k)}{r}"
            if k < len(vals):
                auto(ws, c, value(vals[k]), fmt="0",
                     align=Alignment(horizontal="center", vertical="center", shrink_to_fit=True))
            else:
                inp(ws, c, align=CENTER)
        out = f"D{r}"
        for k in range(5):
            out = f'SUBSTITUTE({out},"{{{k + 1}}}",{get_column_letter(5 + k)}{r})'
        ws[f"B{r}"] = f'=IF(D{r}="","",{out})'
        style(ws, f"B{r}", f=fill("FFFFFF"), fnt=font(11, color="111111"),
              align=Alignment(horizontal="left", vertical="center", wrap_text=False))
    ws.freeze_panes = f"A{R_FIRST}"
    protect(ws)
    return ws


def build_types(wb):
    ws = wb.create_sheet("За типами цілей")
    ws.sheet_properties.tabColor = C_NAVY
    page(ws)
    widths(ws, {"A": 5, "B": 26, **{get_column_letter(c): 11 for c in range(3, 19)}})
    title(ws, "A1:R1", f'="ЦІЛІ ЗА ТИПАМИ — "&{q(REF)}!$B$2&", "&'
                       f'IF({q(SUM_SHEET)}!$C$3="","дата не вказана",'
                       f'RIGHT("0"&DAY({q(SUM_SHEET)}!$C$3),2)&"."&RIGHT("0"&MONTH({q(SUM_SHEET)}!$C$3),2)&"."&'
                       f'YEAR({q(SUM_SHEET)}!$C$3))')
    ws.row_dimensions[1].height = 28

    nb = len(CALC_BLOCKS)                         # C.. — блоки «Розрахунку», далі ефективність
    cols = [get_column_letter(3 + b) for b in range(nb)]
    c_app, c_des, c_eff = cols[1], cols[-1], get_column_letter(3 + nb)
    section(ws, f"A3:{c_eff}3", "УСЬОГО ЗА ТИПАМИ ЦІЛЕЙ")
    heads = ["№", "Тип цілі", "Виявлено", "Застосовано", *[f"Знищено: {lab}" for _, lab in GROUPS],
             "Знищено всього", "Ефектив-ність"]
    for i, t in enumerate(heads):
        header(ws, f"{get_column_letter(1 + i)}4", t)
    ws.row_dimensions[4].height = 42
    tot = T_LAST + 1
    for t in range(TYPE_SLOTS):
        r = T_FIRST + t
        auto(ws, f"A{r}", f'=IF(B{r}="","",{t + 1})')
        auto(ws, f"B{r}", f"={q(CALC)}!$A{r}", align=LEFT)
        for b, col in enumerate(cols):
            auto(ws, f"{col}{r}", f'=IF($B{r}="","",{q(CALC)}!{calc_col(b)}{r})', fmt=NUM, bold=col == c_des)
        auto(ws, f"{c_eff}{r}", f'=IF($B{r}="","",IF(N({c_app}{r})=0,"–",{c_des}{r}/{c_app}{r}))', fmt=PCT)
    label(ws, f"A{tot}:B{tot}", "РАЗОМ (без патрулювання)")
    auto(ws, f"C{tot}", f'=SUM(C{T_FIRST}:C{T_LAST})-SUMIF($B${T_FIRST}:$B${T_LAST},"{PATROL}",'
                        f'C{T_FIRST}:C{T_LAST})', fmt=NUM, bold=True)
    for col in cols[1:]:
        auto(ws, f"{col}{tot}", f"=SUM({col}{T_FIRST}:{col}{T_LAST})", fmt=NUM, bold=True)
    auto(ws, f"{c_eff}{tot}", f'=IF({c_app}{tot}=0,"–",{c_des}{tot}/{c_app}{tot})', fmt=PCT, bold=True)
    style(ws, f"A{tot}:{c_eff}{tot}", f=fill("DDE3EA"))
    ws.conditional_formatting.add(f"{c_des}{T_FIRST}:{c_des}{T_LAST}",
                                  DataBarRule(start_type="num", start_value=0, end_type="max",
                                              color="7FB3A8", showValue=True))

    m_head = tot + 3
    m_first = m_head + 2
    section(ws, f"A{m_head}:R{m_head}", "ЗНИЩЕНО: ТИП ЦІЛІ × ПІДРОЗДІЛ")
    h = m_head + 1
    header(ws, f"A{h}", "№")
    header(ws, f"B{h}", "Тип цілі")
    for u in range(UNIT_SLOTS):
        col = get_column_letter(3 + u)
        header(ws, f"{col}{h}", f"={q(CALC)}!{calc_col(0, u)}$4")
    header(ws, f"R{h}", "Разом")
    for t in range(TYPE_SLOTS):
        r, cr = m_first + t, T_FIRST + t
        auto(ws, f"A{r}", f'=IF(B{r}="","",{t + 1})')
        auto(ws, f"B{r}", f"={q(CALC)}!$A{cr}", align=LEFT)
        for u in range(UNIT_SLOTS):
            col = get_column_letter(3 + u)
            auto(ws, f"{col}{r}", f'=IF(OR($B{r}="",{col}${h}=""),"",{q(CALC)}!{calc_col(nb - 1, u)}{cr})',
                 fmt=NUM)
        auto(ws, f"R{r}", f'=IF($B{r}="","",SUM(C{r}:Q{r}))', fmt=NUM, bold=True)
    mt = m_first + TYPE_SLOTS
    label(ws, f"A{mt}:B{mt}", "РАЗОМ")
    for c in range(3, 19):
        col = get_column_letter(c)
        auto(ws, f"{col}{mt}", f"=SUM({col}{m_first}:{col}{mt - 1})", fmt=NUM, bold=True)
    style(ws, f"A{mt}:R{mt}", f=fill("DDE3EA"))
    ws.freeze_panes = "C5"
    protect(ws)
    return ws


# ================================================================ Картки
CARDS = "Картки"
# Поля картки: ключ, назва, значення за замовчуванням (None — поле без замовчування), список для меню
CARD_FIELDS = [
    ("unit", "Підрозділ", "", "lst_unit_names"),
    ("kind", "Тип", "РЛС", "lst_objects"),
    ("name", "Назва", None, None),
    ("qty", "Кількість", 1, None),
    ("ready", "Боєздатність", "боєздатний", "lst_readiness"),
    ("affil", "Приналежність", "Дружній", "lst_affil"),
    ("source", "Джерело", "Ручна робота", None),
    ("seen", "Час виявлення", "сьогодні", None),
    ("coord", "Координати (MGRS або WGS84)", None, None),
]
CARD_ROW = {key: 5 + i for i, (key, *_) in enumerate(CARD_FIELDS)}
CARD_DEFAULT = {key: default for key, _, default, _ in CARD_FIELDS}
CARD_CHECK = 5 + len(CARD_FIELDS)
CARD_LABEL = CARD_CHECK + 3                  # підпис на мапі, під ним — текст картки
CARD_LINES = [("Тип: ", "kind"), ("Назва: ", "name"), ("Кількість: ", "qty"), ("Боєздатність: ", "ready"),
              ("Приналежність: ", "affil"), ("Підрозділ: ", "unit"), ("Джерело: ", "source"),
              ("Час виявлення: ", "seen"), ("MGRS: ", "mgrs"), ("WGS84: ", "wgs"), ("", "url")]
CARD_TEXT = range(CARD_LABEL + 1, CARD_LABEL + 1 + len(CARD_LINES))
# Службові рядки (приховані): ключ і підпис
CARD_CALC = [
    ("active", "картка заповнена"), ("s", "очищений текст координат"),
    ("m", "MGRS без пробілів"), ("zl", "цифр у номері зони"), ("zone", "зона"), ("band", "смуга"),
    ("col", "літера стовпця 100 км"), ("row", "літера рядка 100 км"), ("dg", "цифри"),
    ("half", "цифр на координату"), ("bpos", "№ смуги"), ("cpos", "№ стовпця 100 км"), ("rpos", "№ рядка 100 км"),
    ("is_m", "це MGRS"), ("e1", "easting"), ("n100", "northing 100 км (без циклу)"), ("n1", "northing"),
    ("xi", "ξ"), ("eta", "η"), ("xp", "ξ'"), ("ep", "η'"), ("chi", "χ"), ("lat_m", "широта з MGRS"),
    ("lon_m", "довгота з MGRS"),
    ("w", "текст WGS84"), ("sep", "позиція роздільника"),
    ("la_s", "широта: текст"), ("la_u", "без знака"), ("la_i", "ціла частина"), ("la_f", "дробова частина"),
    ("la_ok", "широта розпізнана"), ("la", "широта"),
    ("lo_s", "довгота: текст"), ("lo_u", "без знака"), ("lo_i", "ціла частина"), ("lo_f", "дробова частина"),
    ("lo_ok", "довгота розпізнана"), ("lo", "довгота"),
    ("is_w", "це WGS84"), ("lat", "широта"), ("lon", "довгота"),
    ("zone2", "зона UTM"), ("dl", "Δλ"), ("t", "t"), ("xi2", "ξ'"), ("eta2", "η'"), ("e2", "easting"),
    ("n2", "northing"), ("mgrs_w", "MGRS з WGS84"),
    ("mgrs", "MGRS"), ("wgs", "WGS84"), ("url", "посилання"),
]
CALC_FIRST = CARD_TEXT[-1] + 3

# Еліпсоїд WGS84 і проєкція UTM: ряди Крюгера до n³ (похибка менше 1 мм у межах зони)
_FLAT = 1 / 298.257223563
_N = _FLAT / (2 - _FLAT)
UTM_KA = 0.9996 * 6378137 / (1 + _N) * (1 + _N ** 2 / 4 + _N ** 4 / 64)
UTM_E = 2 * math.sqrt(_N) / (1 + _N)
UTM_ALPHA = [_N / 2 - 2 * _N ** 2 / 3 + 5 * _N ** 3 / 16, 13 * _N ** 2 / 48 - 3 * _N ** 3 / 5, 61 * _N ** 3 / 240]
UTM_BETA = [_N / 2 - 2 * _N ** 2 / 3 + 37 * _N ** 3 / 96, _N ** 2 / 48 + _N ** 3 / 15, 17 * _N ** 3 / 480]
UTM_DELTA = [2 * _N - 2 * _N ** 2 / 3 - 2 * _N ** 3, 7 * _N ** 2 / 3 - 8 * _N ** 3 / 5, 56 * _N ** 3 / 15]
MGRS_BANDS = "CDEFGHJKLMNPQRSTUVWX"
MGRS_ROWS = "ABCDEFGHJKLMNPQRSTUV"
MGRS_SETS = '"ABCDEFGH","JKLMNPQR","STUVWXYZ"'
# найменший northing кожної смуги (для вибору 2000-км циклу літер рядка)
MGRS_MIN_N = [1100000, 2000000, 2800000, 3700000, 4600000, 5500000, 6400000, 7300000, 8200000, 9100000,
              0, 800000, 1700000, 2600000, 3500000, 4400000, 5300000, 6200000, 7000000, 7900000]


def num(x):
    return repr(x).upper()


def card_col(i):
    return get_column_letter(4 + 2 * i)


def digit_count(t, n):
    """Скільки з перших n символів t — цифри."""
    return f"SUMPRODUCT(--ISNUMBER(--MID({t},ROW($1:${n}),1)))"


def fixed(x, places, sep):
    """Число з places знаками після роздільника sep — без TEXT(), щоб не залежати від мови Excel."""
    k = 10 ** places
    r = f"ROUND(ABS({x})*{k},0)"
    return f'IF({x}<0,"-","")&INT({r}/{k})&"{sep}"&RIGHT("{"0" * places}"&MOD({r},{k}),{places})'


def card_formulas(c):
    """Службові формули картки у стовпці c: розбір координат, MGRS ↔ WGS84, посилання на мапу."""
    row = {key: CALC_FIRST + i for i, (key, _) in enumerate(CARD_CALC)}
    H = lambda key: f"{c}{row[key]}"
    X = f"{c}{CARD_ROW['coord']}"
    kA = num(UTM_KA)
    sums = lambda coef, f1, x, f2, y: "+".join(f"{num(k)}*{f1}({2 * j}*{x})*{f2}({2 * j}*{y})"
                                               for j, k in enumerate(coef, 1))
    f = {"active": f"=COUNTA({c}{CARD_ROW['unit']}:{X})>0"}

    # очищення: посилання Google Maps, нерозривні пробіли, переноси, «°», дужки
    s = f'UPPER(IF(ISNUMBER(SEARCH("query=",{X})),MID({X},SEARCH("query=",{X})+6,99),{X}))'
    for a, b in [("CHAR(160)", '" "'), ("CHAR(9)", '" "'), ("CHAR(10)", '" "'), ('"°"', '""'),
                 ('"%2C"', '","'), ('"("', '""'), ('")"', '""')]:
        s = f"SUBSTITUTE({s},{a},{b})"
    f["s"] = f"=TRIM({s})"

    # MGRS: «37U CR 12345 67890», пробіли не обов'язкові, кириличні двійники літер замінюються
    m = f'SUBSTITUTE({H("s")}," ","")'
    for cyr, lat in zip("АВСЕНКМРТХУ", "ABCEHKMPTXU"):
        m = f'SUBSTITUTE({m},"{cyr}","{lat}")'
    f["m"] = "=" + m
    M, zl = H("m"), H("zl")
    f["zl"] = f"=IF(ISNUMBER(--MID({M},2,1)),2,1)"
    f["zone"] = f"=IFERROR(--LEFT({M},{zl}),0)"
    f["band"] = f"=MID({M},{zl}+1,1)"
    f["col"] = f"=MID({M},{zl}+2,1)"
    f["row"] = f"=MID({M},{zl}+3,1)"
    f["dg"] = f"=MID({M},{zl}+4,20)"
    f["half"] = f'=LEN({H("dg")})/2'
    zone, dg, half = H("zone"), H("dg"), H("half")
    f["bpos"] = f'=IF({H("band")}="",0,IFERROR(FIND({H("band")},"{MGRS_BANDS}"),0))'
    f["cpos"] = f'=IF({H("col")}="",0,IFERROR(FIND({H("col")},CHOOSE(MOD({zone}-1,3)+1,{MGRS_SETS})),0))'
    f["rpos"] = f'=IF({H("row")}="",0,IFERROR(FIND({H("row")},"{MGRS_ROWS}"),0))'
    f["is_m"] = (f'=AND({zone}>=1,{zone}<=60,{H("bpos")}>0,{H("cpos")}>0,{H("rpos")}>0,{half}>=1,{half}<=5,'
                 f"{half}=INT({half}),{digit_count(dg, 10)}=LEN({dg}))")
    f["e1"] = f'={H("cpos")}*100000+IFERROR(--LEFT({dg},{half}),0)*10^(5-{half})'
    f["n100"] = f'=MOD({H("rpos")}-1-IF(MOD({zone},2)=0,5,0),20)*100000'
    min_n = "{" + ",".join(map(str, MGRS_MIN_N)) + "}"
    n100 = H("n100")
    f["n1"] = (f'={n100}+2000000*MAX(0,ROUNDUP((INDEX({min_n},MAX(1,{H("bpos")}))-{n100})/2000000,0))'
               f"+IFERROR(--RIGHT({dg},{half}),0)*10^(5-{half})")
    # UTM → широта/довгота
    f["xi"] = f'=({H("n1")}-IF({H("bpos")}<{MGRS_BANDS.index("N") + 1},10000000,0))/{kA}'
    f["eta"] = f'=({H("e1")}-500000)/{kA}'
    xi, eta = H("xi"), H("eta")
    f["xp"] = f"={xi}-({sums(UTM_BETA, 'SIN', xi, 'COSH', eta)})"
    f["ep"] = f"={eta}-({sums(UTM_BETA, 'COS', xi, 'SINH', eta)})"
    f["chi"] = f'=ASIN(SIN({H("xp")})/COSH({H("ep")}))'
    chi = H("chi")
    f["lat_m"] = f"=DEGREES({chi}+" + "+".join(f"{num(d)}*SIN({2 * j}*{chi})" for j, d in enumerate(UTM_DELTA, 1)) + ")"
    f["lon_m"] = f'=({zone}-1)*6-177+DEGREES(ATAN(SINH({H("ep")})/COS({H("xp")})))'

    # WGS84: «50.12345, 30.12345», «50,12345 30,12345», «50.1;30.1», «N50.1 E30.1»
    f["w"] = f'=TRIM(SUBSTITUTE(SUBSTITUTE({H("s")},"N",""),"E",""))'
    W = H("w")
    commas = f'(LEN({W})-LEN(SUBSTITUTE({W},",","")))'
    f["sep"] = (f'=IF(ISNUMBER(FIND(";",{W})),FIND(";",{W}),IF(ISNUMBER(FIND(", ",{W})),FIND(", ",{W}),'
                f'IF(ISNUMBER(FIND(" ",{W})),FIND(" ",{W}),'
                f'IFERROR(FIND("|",SUBSTITUTE({W},",","|",ROUNDUP({commas}/2,0))),0))))')
    sep = H("sep")
    for k, part, n_int in [("la", f"LEFT({W},{sep}-1)", 2), ("lo", f"MID({W},{sep}+1,40)", 3)]:
        f[f"{k}_s"] = f'=IF({sep}=0,"",SUBSTITUTE(TRIM({part}),",","."))'
        t = H(f"{k}_s")
        f[f"{k}_u"] = f'=IF(OR(LEFT({t},1)="-",LEFT({t},1)="+"),MID({t},2,40),{t})'
        u = H(f"{k}_u")
        f[f"{k}_i"] = f'=LEFT({u},FIND(".",{u}&".")-1)'
        f[f"{k}_f"] = f'=MID({u},FIND(".",{u}&".")+1,40)'
        i, fr = H(f"{k}_i"), H(f"{k}_f")
        f[f"{k}_ok"] = (f"=AND(LEN({i})>=1,LEN({i})<={n_int},{digit_count(i, n_int)}=LEN({i}),"
                        f"LEN({fr})<=15,{digit_count(fr, 15)}=LEN({fr}))")
        f[k] = f'=IF({H(k + "_ok")},IF(LEFT({t},1)="-",-1,1)*(--{i}+IF({fr}="",0,--{fr}/10^LEN({fr}))),0)'
    la, lo = H("la"), H("lo")
    f["is_w"] = (f'=AND(NOT({H("is_m")}),{H("la_ok")},{H("lo_ok")},{la}>=-80,{la}<=84,'
                 f"{lo}>=-180,{lo}<=180)")
    f["lat"] = f'=IF({H("is_m")},{H("lat_m")},IF({H("is_w")},{la},0))'
    f["lon"] = f'=IF({H("is_m")},{H("lon_m")},IF({H("is_w")},{lo},0))'

    # широта/довгота → UTM → MGRS (з винятками зон для Норвегії та Шпіцбергена)
    lat, lon = H("lat"), H("lon")
    f["zone2"] = (f"=IF(AND({lat}>=56,{lat}<64,{lon}>=3,{lon}<12),32,IF(AND({lat}>=72,{lon}>=0,{lon}<42),"
                  f"IF({lon}<9,31,IF({lon}<21,33,IF({lon}<33,35,37))),MIN(60,INT(({lon}+180)/6)+1)))")
    z2 = H("zone2")
    f["dl"] = f"=RADIANS({lon}-({z2}-1)*6+177)"
    sp = f"SIN(RADIANS({lat}))"
    f["t"] = f"=SINH(ATANH({sp})-{num(UTM_E)}*ATANH({num(UTM_E)}*{sp}))"
    f["xi2"] = f'=ATAN({H("t")}/COS({H("dl")}))'
    f["eta2"] = f'=ATANH(SIN({H("dl")})/SQRT(1+{H("t")}^2))'
    xi2, eta2 = H("xi2"), H("eta2")
    f["e2"] = f"=500000+{kA}*({eta2}+{sums(UTM_ALPHA, 'COS', xi2, 'SINH', eta2)})"
    f["n2"] = f"={kA}*({xi2}+{sums(UTM_ALPHA, 'SIN', xi2, 'COSH', eta2)})+IF({lat}<0,10000000,0)"
    e2, n2 = H("e2"), H("n2")
    pad = lambda x: f'RIGHT("0000"&MOD(INT({x}),100000),5)'
    f["mgrs_w"] = (f'={z2}&MID("{MGRS_BANDS}X",INT(({lat}+80)/8)+1,1)&" "&'
                   f"MID(CHOOSE(MOD({z2}-1,3)+1,{MGRS_SETS}),INT({e2}/100000),1)&"
                   f'MID("{MGRS_ROWS}",MOD(INT({n2}/100000)+IF(MOD({z2},2)=0,5,0),20)+1,1)&" "&'
                   f'{pad(e2)}&" "&{pad(n2)}')

    # результат
    is_m, is_w = H("is_m"), H("is_w")
    f["mgrs"] = (f'=IF({is_m},{zone}&{H("band")}&" "&{H("col")}&{H("row")}&" "&LEFT({dg},{half})&" "&'
                 f'RIGHT({dg},{half}),IF({is_w},{H("mgrs_w")},""))')
    ok = f"OR({is_m},{is_w})"
    f["wgs"] = f'=IF({ok},{fixed(lat, 6, ",")}&", "&{fixed(lon, 6, ",")},"")'
    f["url"] = (f'=IF({ok},"https://www.google.com/maps/search/?api=1&query="&{fixed(lat, 7, ".")}&"%2C"&'
                f'{fixed(lon, 7, ".")},"")')
    return f, H


def build_cards(wb):
    ws = wb.create_sheet(CARDS)
    ws.sheet_properties.tabColor = C_TEAL
    page(ws)
    last = card_col(CARD_SLOTS - 1)
    widths(ws, {"A": 26, "B": 22, "C": 2})
    for i in range(CARD_SLOTS):
        ws.column_dimensions[card_col(i)].width = 44
        ws.column_dimensions[get_column_letter(5 + 2 * i)].width = 2
    title(ws, f"A1:{last}1", "КАРТКИ ОБ'ЄКТІВ — текст картки для месенджера складається автоматично", size=14)
    ws.row_dimensions[1].height = 34
    ws["A2"] = ("Одна картка — один стовпець. Заповніть жовті клітинки; порожні поля беруться зі стовпця "
                "«За замовчуванням». Координати — MGRS (37U AB 12345 67890) або WGS84 (50.12345, 30.12345): "
                "друге представлення і посилання на мапу рахуються самі.")
    style(ws, f"A2:{card_col(2)}2", fnt=font(9, italic=True, color="5B6B7F"), border=None, align=LEFT, merge=True)
    ws.row_dimensions[2].height = 30

    # --- поля
    section(ws, f"A3:{last}3", "ДАНІ ОБ'ЄКТА")
    header(ws, "A4", "Поле")
    header(ws, "B4", "За замовчуванням")
    for i in range(CARD_SLOTS):
        header(ws, f"{card_col(i)}4", f"Картка {i + 1}")
    cols = [card_col(i) for i in range(CARD_SLOTS)]
    for key, name, default, lst in CARD_FIELDS:
        r = CARD_ROW[key]
        label(ws, f"A{r}", name)
        if default is None:
            auto(ws, f"B{r}", "—")
        else:
            inp(ws, f"B{r}", default if default != "" else None, align=CENTER)
        fmt = "0" if key == "qty" else "@"
        for c in cols:
            inp(ws, f"{c}{r}", fmt=fmt)
        if lst:
            list_dv(ws, lst, [f"B{r}", f"{cols[0]}{r}:{last}{r}"], strict=False)
    ws[f"B{CARD_ROW['unit']}"].comment = Comment("Наприклад, назва підрозділу так, як її пишуть у картках.", "ППО")
    ws[f"A{CARD_ROW['coord']}"].comment = Comment(
        "MGRS: 37U AB 12345 67890 (можна без пробілів, точність 1–5 цифр).\n"
        "WGS84: 50.12345, 30.12345 або 50,12345 30,12345.\n"
        "Можна вставити посилання Google Maps з ?query=.", "ППО")
    num_dv(ws, [f"B{CARD_ROW['qty']}", f"{cols[0]}{CARD_ROW['qty']}:{last}{CARD_ROW['qty']}"], hi=9999)
    label(ws, f"A{CARD_CHECK}", "Перевірка")
    auto(ws, f"B{CARD_CHECK}")
    ws.row_dimensions[CARD_ROW["coord"]].height = 20

    # --- текст картки
    section(ws, f"A{CARD_LABEL - 1}:{last}{CARD_LABEL - 1}",
            f"ТЕКСТ КАРТКИ — виділіть рядки {CARD_TEXT[0]}–{CARD_TEXT[-1]} потрібної картки → Ctrl+C → "
            "вставте в месенджер")
    label(ws, f"A{CARD_LABEL}:B{CARD_LABEL}", "Підпис на мапі (Підрозділ_Назва)")
    label(ws, f"A{CARD_TEXT[0]}:B{CARD_TEXT[-1]}", "Текст картки (копіюйте ці рядки)")
    ws[f"A{CARD_TEXT[0]}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[CARD_LABEL].height = 22

    def eff(c, key):
        """Значення поля з урахуванням стовпця «За замовчуванням»."""
        r = CARD_ROW[key]
        if CARD_DEFAULT[key] is None:
            return f'{c}{r}&""'
        return f'IF({c}{r}="",$B${r},{c}{r})&""'

    first_calc, last_calc = CALC_FIRST, CALC_FIRST + len(CARD_CALC) - 1
    for i, (key, lab) in enumerate(CARD_CALC):
        ws[f"A{first_calc + i}"] = lab
        ws[f"A{first_calc + i}"].font = font(8, italic=True, color="7F7F7F")
    ws[f"A{first_calc - 1}"] = "СЛУЖБОВІ РОЗРАХУНКИ КООРДИНАТ — не редагуються"
    ws[f"A{first_calc - 1}"].font = font(8, True, "7F7F7F")

    for c in cols:
        f, H = card_formulas(c)
        for key, formula in f.items():
            ws[H(key)] = formula
            ws[H(key)].font = font(8, color="7F7F7F")
        active, X = H("active"), f"{c}{CARD_ROW['coord']}"
        auto(ws, f"{c}{CARD_CHECK}", f'=IF(NOT({active}),"",IF({X}="","⚠ вкажіть координати",'
                                     f'IF(OR({H("is_m")},{H("is_w")}),"✓","⚠ координати не розпізнано")))')
        unit, name = eff(c, "unit"), eff(c, "name")
        ws[f"{c}{CARD_LABEL}"] = f'=IF({active},IF({unit}="",{name},{unit}&"_"&{name}),"")'
        style(ws, f"{c}{CARD_LABEL}", f=fill(C_AUTO), fnt=font(11, True, "43BF4D"),
              align=Alignment(horizontal="left", vertical="center", indent=1))
        for r, (prefix, key) in zip(CARD_TEXT, CARD_LINES):
            if key == "url":
                v = f'IF({H("url")}="","",HYPERLINK({H("url")},{H("url")}))'
            elif key in ("mgrs", "wgs"):
                v = f'"{prefix}"&{H(key)}'
            else:
                v = f'"{prefix}"&{eff(c, key)}'
            ws[f"{c}{r}"] = f'=IF({active},{v},"")'
            edge = Border(left=THIN, right=THIN, bottom=THIN if r == CARD_TEXT[-1] else None)
            style(ws, f"{c}{r}", f=fill("FFFFFF"), border=edge,
                  fnt=font(9, color="1F5FAD") if key == "url" else font(10, color="111111"),
                  align=Alignment(horizontal="left", vertical="center", indent=1, shrink_to_fit=key == "url"))

    # картка заповнена → підпис у кольорах мапи; позначка перевірки
    lab_rng = f"{cols[0]}{CARD_LABEL}:{last}{CARD_LABEL}"
    ws.conditional_formatting.add(lab_rng, FormulaRule(formula=[f'{cols[0]}{CARD_LABEL}<>""'], fill=fill("2B363F")))
    chk = f"{cols[0]}{CARD_CHECK}:{last}{CARD_CHECK}"
    ws.conditional_formatting.add(chk, FormulaRule(formula=[f'LEFT({cols[0]}{CARD_CHECK},1)="⚠"'],
                                                   fill=fill("FDE2C4"), font=Font(name=FONT, color="9A3412", bold=True)))
    ws.conditional_formatting.add(chk, CellIsRule(operator="equal", formula=['"✓"'],
                                                  font=Font(name=FONT, color="1E6B34", bold=True)))

    ws.row_dimensions.group(first_calc - 1, last_calc, hidden=True)
    ws.freeze_panes = "C5"
    ws.print_area = f"A1:{last}{CARD_TEXT[-1]}"
    ws.print_title_cols = "A:B"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 0, 1
    protect(ws)
    return ws


# ================================================================ Інструкція
def build_help(wb, cfg):
    ws = wb.active
    ws.title = "Інструкція"
    ws.sheet_properties.tabColor = C_NAVY
    page(ws, landscape=False)
    widths(ws, {"A": 3, "B": 24, "C": 90})
    title(ws, "A1:C1", "ДОБОВА ДОПОВІДЬ ППО — як користуватися", size=16)
    ws.row_dimensions[1].height = 50
    emblem(ws, cfg.get("org_emblem"), 2, 0, 60, 60, off_x=560, off_y=3)
    r = 3

    def head(text):
        nonlocal r
        section(ws, f"B{r}:C{r}", text)
        r += 1

    def line(a, b, f_a=None):
        nonlocal r
        ws[f"B{r}"], ws[f"C{r}"] = a, b
        style(ws, f"B{r}", f=f_a, fnt=font(10, True, C_SLATE), align=LEFT_TOP, border=None)
        style(ws, f"C{r}", fnt=font(10), align=LEFT_TOP, border=None)
        ws.row_dimensions[r].height = max(16, 15 * (1 + len(b) // 95))
        r += 1

    head("ЩОДНЯ")
    line("1. Дата", "На аркуші «Зведення» вкажіть дату звіту (клітинка C3). Вона підтягнеться на всі аркуші.")
    line("2. Доповіді", "Скопіюйте готовий текст доповіді (з месенджера чи документа) і вставте у стовпець B "
                        "«Текст доповіді» журналу свого підрозділу — одна доповідь в один рядок.")
    line("3. Розбір", "Час, тип цілі, район, квадрат, засіб ураження, боєприпас, екіпаж, результат і витрати "
                      "заповнюються з тексту автоматично (блакитні клітинки).")
    line("Засіб і боєприпас", "Засіб ураження розпізнається за назвою моделі чи системи або за шифром "
                              "боєприпасу. Якщо модель не вказана — ставиться загальний засіб (FPV-перехоплювач, ПЗРК). "
                              "Боєприпас: якщо в засобу він один — підставляється сам; якщо кілька — шукається в "
                              "тексті, інакше оберіть зі списку (у списку лише боєприпаси обраного засобу).")
    line("4. Перевірка", "Стовпець «Перевірка» показує ✓ або ⚠. Якщо тип цілі чи результат не розпізнано, "
                         "клітинка підсвічується помаранчевим — оберіть правильне значення зі списку прямо в ній.")
    line("5. Удари противника", "Тексти про РУ / АУ / УДК вставляйте у стовпець R праворуч від журналу.")
    line("6. Зведення", "«Зведення», «За типами цілей» і «Розрахунок» рахуються самі; вручну там нічого "
                        "не вводиться.")
    line("Доповідь", "Аркуш «Доповідь» складає текст доповіді у встановленому форматі. Оберіть «Усі підрозділи» "
                     "або конкретний підрозділ, виділіть стовпець B і скопіюйте. Формулювання рядків і перелік "
                     "типів для кожного показника можна змінити там же (жовті клітинки).")
    line("Картки", "Аркуш «Картки» складає текст картки об'єкта (РЛС, ЗРК, екіпаж тощо): тип, назва, кількість, "
                   "боєздатність, приналежність, підрозділ, джерело, час виявлення, MGRS, WGS84 і посилання на "
                   "Google Maps. Детальніше — розділ «Картки об'єктів» нижче.")
    line("Групи засобів", "Зведення ділить засоби ураження на групи: FPV-перехоплювачі, стрілецька зброя, "
                          "зенітна артилерія, ЗРК/ПЗРК, інші. Група кожного засобу задається в «Довідниках».")
    line("7. Новий день", "Збережіть файл як копію з датою в назві. Щоб очистити журнал, виділяйте ТІЛЬКИ "
                          "стовпці з текстом (B і R) та жовті поля й натискайте Delete — блакитні клітинки "
                          "містять формули розбору, їх не стирайте.")
    r += 1
    head("ПОРАДИ ДЛЯ ВСТАВКИ")
    line("Одна клітинка", "Якщо в доповіді є переноси рядків, двічі клацніть на клітинку (або F2) і вставте "
                          "текст усередину — тоді він не розійдеться на кілька рядків.")
    line("Що розпізнається", "Час — перші символи (07:24 або 07.24). Квадрат — «(37U XX 12345 67890)». "
                             "Район — після «н.п.», «в районі», «в р-ні» або перед назвою області / «кв.». "
                             "Тип цілі — після «виявлено» або «Ціль:». Витрати — «Витрата: 1 борт», "
                             "втрати — «Втрати: 1 борт» або «Борт втрачено».")
    line("Нові слова", "Якщо якийсь тип цілі чи засіб постійно не розпізнається — допишіть ключове слово "
                       "в таблицю «Ключові слова» на аркуші «Довідники». Новий засіб ураження спершу додайте до "
                       "списку засобів (з групою), а його боєприпаси — до таблиці «Засіб — Боєприпас».")
    line("Виправлення", "Вибране вручну значення замінює формулу лише в цьому рядку. Щоб повернути автоматичний "
                        "розбір, скопіюйте клітинку з сусіднього порожнього рядка.")
    r += 1
    head("КАРТКИ ОБ'ЄКТІВ")
    line("Заповнення", "Одна картка — один стовпець аркуша «Картки». Заповніть жовті клітинки; порожні поля "
                       "беруться зі стовпця «За замовчуванням» (його теж можна змінити: тип, кількість, "
                       "боєздатність, приналежність, джерело, час). Підрозділ достатньо вказати там один раз.")
    line("Координати", "Вводьте MGRS (37U AB 12345 67890, можна без пробілів і з точністю 1–5 цифр) або WGS84 "
                       "(50.12345, 30.12345 чи 50,12345 30,12345); можна вставити й посилання Google Maps. "
                       "Друге представлення і посилання на мапу рахуються самі, «⚠» — координати не розпізнано.")
    line("Копіювання", f"Виділіть рядки {CARD_TEXT[0]}–{CARD_TEXT[-1]} потрібної картки (від «Тип» до посилання) "
                       "→ Ctrl+C → вставте в месенджер. Над текстом — підпис для мапи «Підрозділ_Назва».")
    line("Списки", "Типи об'єктів, варіанти боєздатності й приналежності — на аркуші «Довідники» праворуч. "
                   "Значення поза списком теж можна ввести (Excel лише попередить).")
    r += 1
    head("КОЛЬОРИ")
    line("Жовті клітинки", "заповнюються вручну (текст доповіді, черговий, РЛС, екіпажі, картки)", f_a=fill(C_INPUT))
    line("Блакитні клітинки", "розібрано з тексту автоматично; можна виправити вручну", f_a=fill(C_PARSED))
    line("Сірі клітинки", "рахуються автоматично, захищені від змін", f_a=fill(C_AUTO))
    line("Рядок «Приклад»", "показує формат; у підсумках не враховується", f_a=fill(C_EXAMPLE))
    r += 1
    head("ПРАВА ДОСТУПУ")
    line("Зараз", "Усі аркуші захищені від випадкових змін БЕЗ пароля: редагувати можна лише жовті й блакитні клітинки.")
    line("Пароль підрозділу", "Рецензування → Зняти захист аркуша → Захистити аркуш → задайте окремий пароль "
                              "для кожного аркуша підрозділу й передайте його лише відповідальному.")
    line("Адміністратор", "Так само поставте свій пароль на «Зведення», «За типами цілей», «Довідники», "
                          "«Розрахунок» і на структуру книги (Рецензування → Захистити книгу).")
    line("Важливо", "Пароль аркуша в Excel захищає від помилок, але не є надійним захистом даних. "
                    "Для справжнього розмежування доступу зберігайте файл у OneDrive/SharePoint із правами "
                    "на рівні файлу або ведіть окремий файл для кожного підрозділу.")
    r += 1
    head("ЯК ДОДАТИ ПІДРОЗДІЛ")
    line("1", "Рецензування → Захистити книгу (зняти захист структури).")
    line("2", "Правою кнопкою на вкладці будь-якого підрозділу → Перемістити або копіювати → «Створити копію».")
    line("3", "Перейменуйте вкладку на новий код (наприклад, «П4»), очистіть жовті клітинки й замініть шеврон "
              "(клацніть на зображення → Змінити рисунок).")
    line("4", "На аркуші «Довідники» допишіть код (точно як назва вкладки) і повну назву в перший "
              "вільний рядок. Зведення підхопить підрозділ автоматично (до 15 підрозділів).")
    line("Перевірка", "У «Зведенні» стовпець «Аркуш» показує ✓, якщо код у довіднику збігається з вкладкою.")
    ws.protection.sheet = True
    return ws


# ================================================================ main
def build(cfg):
    REGION_WORDS[:] = cfg.get("region_words", [])
    if len(cfg["units"]) > UNIT_SLOTS:
        raise SystemExit(f"Не більше {UNIT_SLOTS} підрозділів")
    wb = Workbook()
    build_help(wb, cfg)
    build_summary(wb, cfg)
    build_report(wb)
    build_cards(wb)
    build_types(wb)
    for i, unit in enumerate(cfg["units"], 1):
        build_unit(wb, unit, i, cfg.get("radars", []))
    build_reference(wb, cfg)
    build_calc(wb)
    wb.security = WorkbookProtection(lockStructure=True)
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    return wb


def main():
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as fh:
        cfg = json.load(fh)
    base = os.path.dirname(os.path.abspath(sys.argv[1]))
    resolve = lambda p: p if not p or os.path.isabs(p) else os.path.join(base, p)
    cfg["org_emblem"] = resolve(cfg.get("org_emblem"))
    for u in cfg["units"]:
        u["emblem"] = resolve(u.get("emblem"))
    build(cfg).save(sys.argv[2])
    print("Збережено:", sys.argv[2])


if __name__ == "__main__":
    main()
